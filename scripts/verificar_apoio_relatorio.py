"""Verificador da v4 (apoio da reunião), §6 do prompt_apoio_reuniao_v4.

1. Os 140 slides da v3 estão na v4, na mesma ordem e idênticos (texto e células de tabela).
2. Cada original tem exatamente um apoio logo depois, com o mesmo ID; P00 abre e G fecha.
3. Cada apoio tem de 2 a 4 perguntas, todas com "De onde"; o que está no slide de apoio está no json.
4. Todo número do apoio (o que é, respostas, de cor) existe no slide original, nos slides de `ver` ou
   na lista de constantes de método (limiares e escalas citados nas leituras).
5. Todo termo técnico usado nos apoios tem verbete no glossário.
6. Nenhuma resposta escolhe pela orientadora; nenhum termo do vocabulário proibido.
7. Nenhuma fonte abaixo de 11 pt nos slides de apoio; notas de cada original com o apoio.
8. Lista A CONFERIR e leituras pendentes (informativo).

    uv run --with python-pptx python scripts/verificar_apoio_relatorio.py
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from pptx import Presentation

RAIZ = Path(__file__).resolve().parents[1]
PASTA = RAIZ / 'relatorio_descritivo'
V3 = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v3.pptx'
V4 = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v4_APOIO.pptx'
A = json.loads((PASTA / 'apoio' / 'apoio.json').read_text(encoding='utf-8'))
ID = re.compile(r'^(A|DM|DS|DC|FA|FC|FP|Z|VP)\d{2}(-\d+)?$')
NUM = re.compile(r'(?<![\w-])-?\d[\d.]*(?:,\d+)?')   # não pega o 7 de IVS-7 nem códigos V00495
FAIXA = 'APOIO · NÃO COMPARTILHAR · '
CONST = {'0', '1', '2', '3', '-1', '0,5', '0,6', '0,7', '0,8', '0,70', '0,50', '0,30', '0,40', '0,79', '0,60',
         '1,5', '250', '30', '60', '40', '60,4', '39,6', '2012', '2022', '100', '50'}
TECNICOS = ['KMO', 'MSA', 'Bartlett', 'comunalidade', 'carga', 'Varimax', 'promax', 'Φ', 'matriz padrão',
            'matriz estrutura', 'Horn', 'Kaiser', 'autovalor', 'AUC', 'ROC', 'Spearman', 'Pearson', 'FCU',
            'Dados_sig', 'sigilo', 'exclusão por lista', 'assimetria', 'min-max', 'normaliza', 'parcimônia',
            'estabilidade', 'amostra comum', 'postos', 'formativo', 'falácia ecológica', 'quartis', 'P10']
PROIBIDAS = ['vale destacar', 'é importante notar', 'em suma', 'robusto', 'abrangente', 'vale ressaltar']
ok_geral = []


def item(n, ok, msg):
    ok_geral.append(ok)
    print(f"{'OK   ' if ok else 'FALHA'} {n}. {msg}")


def conteudo(s):
    txt = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    tab = [[c.text for c in r.cells] for sh in s.shapes if sh.has_table for r in sh.table.rows]
    return txt, tab


def sid(s):
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if ID.match(t):
                return t
            if t.startswith(FAIXA):
                return 'APOIO:' + t[len(FAIXA):]
    return '?'


def nums(t):
    return {n.rstrip('.,') for n in NUM.findall(t)}


def main() -> None:
    v3, v4 = list(Presentation(V3).slides), list(Presentation(V4).slides)
    ids3 = [sid(s) for s in v3]
    ids4 = [sid(s) for s in v4]
    orig4 = [(i, s) for i, s in enumerate(v4) if not ids4[i].startswith('APOIO:')]
    # 1
    iguais = [ids4[i] for i, _ in orig4] == ids3 and all(conteudo(a) == conteudo(b) for a, (_, b) in zip(v3, orig4))
    item(1, iguais, f'{len(orig4)} originais na v4, na ordem e idênticos à v3 (texto e tabelas)')
    # 2
    erros = [ids4[i] for i, _ in orig4 if i + 1 >= len(ids4) or ids4[i + 1] != 'APOIO:' + ids4[i]]
    gs = [x for x in ids4 if re.fullmatch(r'APOIO:G\d{2}', x)]
    ok2 = not erros and ids4[0] == 'APOIO:P00' and ids4[-len(gs):] == gs and len(ids4) == 2 * len(ids3) + 1 + len(gs)
    item(2, ok2, f'cada original seguido do seu apoio; P00 no início; {len(gs)} de glossário no fim'
         + (f'; problemas: {erros[:5]}' if erros else ''))
    # 3
    slides = {k: v for k, v in A.items() if v.get('tipo') in ('slide', 'geral')}
    ruins = [k for k, v in slides.items() if not 2 <= len(v['q']) <= 4 or any(not q['de_onde'] for q in v['q'])]
    fora = []
    for i, s in enumerate(v4):
        if ids4[i].startswith('APOIO:') and not ids4[i].startswith('APOIO:G'):
            k = ids4[i][6:]
            txt = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            fora += [f'{k}: {q["p"][:40]}' for q in A[k]['q'] if q['p'] not in txt]
    item(3, not ruins, f'{len(slides)} apoios com 2 a 4 perguntas e "De onde" em todas'
         + (f'; fora do padrão: {ruins[:5]}' if ruins else '') + f'; perguntas só nas notas (corte): {len(fora)}')
    # 4
    txt3 = {i: ' '.join(t for t in conteudo(s)[0]) + ' ' + ' '.join(' '.join(r) for r in conteudo(s)[1])
            for i, s in zip(ids3, v3)}
    sem = []
    for k, v in slides.items():
        base = set(CONST)
        for x in ([k] if k in txt3 else []) + v.get('ver', []):
            base |= nums(txt3[x])
        texto = ' '.join([v['o'], *[q['r'] for q in v['q']], *[q['p'] for q in v['q']], *v['k']])
        sem += [f'{k}: {n}' for n in sorted(nums(texto) - base)]
    item(4, not sem, 'números dos apoios presentes no slide (ou em `ver`/constantes)'
         + (f'; sem fonte ({len(sem)}): {sem}' if sem else ''))
    # 5
    glos = ' '.join(t['t'] + ' ' + t['d'] for v in A.values() if v.get('tipo') == 'glossario' for t in v['termos']).lower()
    usados = {t for t in TECNICOS for v in slides.values()
              if t.lower() in (v['o'] + ' '.join(q['r'] for q in v['q'])).lower()}
    faltam = sorted(t for t in usados if t.lower() not in glos)
    item(5, not faltam, f'{len(usados)} termos técnicos usados; sem verbete: {faltam or "nenhum"}')
    # 6
    voc = []
    for k, v in slides.items():
        for q in v['q']:
            lo = q['r'].lower()
            if re.search(r'(?<!não )\b(escolhi|a melhor é|melhor versão|a escolha é minha)\b', lo):
                voc.append(f'{k}: escolha')
            voc += [f'{k}: {p}' for p in PROIBIDAS if p in lo]
    item(6, not voc, f'nenhuma escolha pela orientadora e vocabulário limpo: {voc[:5] or "ok"}')
    # 7
    pequenas, notas = [], []
    for i, s in enumerate(v4):
        if ids4[i].startswith('APOIO:'):
            pequenas += [ids4[i] for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs
                         for r in p.runs if r.font.size and r.font.size.pt < 11]
        else:
            nt = ' '.join(sh.text_frame.text for sh in s.notes_slide.shapes if sh.has_text_frame)
            if f'APOIO {ids4[i]}' not in nt:
                notas.append(ids4[i])
    item(7, not pequenas and not notas, f'fonte ≥ 11 pt nos apoios ({len(set(pequenas))} com fonte menor); '
         f'notas dos originais com o apoio ({len(notas)} sem)')
    # 8
    conf = [f'{k}: {q["p"]}' for k, v in slides.items() for q in v['q'] if 'A CONFERIR' in q['de_onde']
            or '[Escreva' in q['r']]
    pend = sorted({p for v in slides.values() for p in v.get('p', [])})
    print(f'INFO  8. A CONFERIR antes da reunião ({len(conf)}):')
    for c in conf:
        print('        ', c)
    print(f'INFO     leituras pendentes ({len(pend)}): ' + ' | '.join(pend))
    print('VERIFICAÇÃO ' + ('VERDE' if all(ok_geral) else 'VERMELHA'))
    sys.exit(0 if all(ok_geral) else 1)


if __name__ == '__main__':
    main()
