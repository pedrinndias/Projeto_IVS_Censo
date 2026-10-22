"""Anexa a parte das 108 versões da fatorial ao relatório descritivo EDITADO pelo Pedro.

O relatório descritivo é editado à mão (slides apagados, textos encurtados) e não pode ser
regerado sem perder essas edições. Este script não regera nada: abre o deck editado, copia
para dentro dele os slides de `versoes_fatorial/dados/parte_relatorio.pptx` (formas,
tabelas, imagens e notas), põe a parte entre o FP13 e o Z02, acrescenta a linha da parte 5 no
roteiro (A03) e renumera o número impresso no pé de todos os slides (os apagados tinham
deixado 74 números fora de lugar). Grava um ARQUIVO NOVO; o deck do Pedro fica intocado.

    uv run --with python-pptx python scripts/anexar_versoes_ao_relatorio.py \
        relatorio_editado.pptx [saida.pptx]
"""
from __future__ import annotations

import copy
import io
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

RAIZ = Path(__file__).resolve().parents[1]
PARTE = RAIZ / 'versoes_fatorial' / 'dados' / 'parte_relatorio.pptx'
SAIDA = RAIZ / 'relatorio_descritivo' / 'Relatorio_Descritivo_IVS_2026-10_com_108_versoes.pptx'
ID = re.compile(r'^(A|DM|DS|DC|FA|FC|FP|Z|VP)\d{2}(-\d+)?$')
R_EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'


def sid(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and ID.match(sh.text_frame.text.strip()):
            return sh.text_frame.text.strip()
    return '?'


def copiar(src, dest_prs):
    """Um slide da parte para dentro do relatório: fundo, formas (com imagens) e notas."""
    novo = dest_prs.slides.add_slide(dest_prs.slides[0].slide_layout)
    for sh in list(novo.shapes):
        sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.bg
    if bg is not None:
        novo._element.cSld.insert(0, copy.deepcopy(bg))
    arvore = novo.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'):
            continue
        el = copy.deepcopy(el)
        for blip in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            parte_img = src.part.related_part(blip.get(R_EMBED))
            _, rid = novo.part.get_or_add_image_part(io.BytesIO(parte_img.blob))
            blip.set(R_EMBED, rid)
        arvore.append(el)
    if src.has_notes_slide:
        texto = src.notes_slide.notes_text_frame.text if src.notes_slide.notes_text_frame else ''
        dest = novo.notes_slide
        if dest.notes_text_frame is not None:
            dest.notes_text_frame.text = texto
        else:   # o modelo de notas do pptxgenjs não traz a caixa de texto: copia as formas da origem
            for el in src.notes_slide.shapes._spTree.iterchildren():
                if not (el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr')):
                    dest.shapes._spTree.append(copy.deepcopy(el))
    return novo


def linha_roteiro(prs, ids_parte):
    """Acrescenta a parte 5 no roteiro (A03), copiando a linha da parte 4."""
    for s in prs.slides:
        if sid(s) != 'A03':
            continue
        for sh in s.shapes:
            if not sh.has_table:
                continue
            tbl = sh.table._tbl
            linhas = tbl.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}tr')
            modelo = next((tr for tr in linhas if tr.xpath('string(.)').strip().startswith('4')), None)
            if modelo is None or any(tr.xpath('string(.)').strip().startswith('5') for tr in linhas):
                return False
            nova = copy.deepcopy(modelo)
            textos = ['5', 'As 108 versões da análise fatorial', f'{ids_parte[0]}–{ids_parte[-1]}',
                      'de 68.000 a 87.545 por versão']
            for tc, t in zip(nova.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}tc'), textos):
                runs = tc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}t')
                runs[0].text = t
                for r in runs[1:]:
                    r.text = ''
            modelo.addnext(nova)
            for sh2 in s.shapes:     # o subtítulo do roteiro contava quatro partes
                if sh2.has_text_frame and 'Quatro partes' in sh2.text_frame.text:
                    for par in sh2.text_frame.paragraphs:
                        for r in par.runs:
                            r.text = r.text.replace('Quatro partes', 'Cinco partes')
            return True
    return False


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit('uso: anexar_versoes_ao_relatorio.py <relatorio_editado.pptx> [saida.pptx]')
    editado = Path(sys.argv[1])
    saida = Path(sys.argv[2]) if len(sys.argv) > 2 else SAIDA
    if saida.resolve() == editado.resolve():
        sys.exit('A saída não pode ser o deck editado: o script grava sempre um arquivo novo.')
    prs, parte = Presentation(editado), Presentation(PARTE)
    antes = [sid(s) for s in prs.slides]
    if any(i.startswith('VP') for i in antes):
        sys.exit('O relatório já tem a parte das 108 versões (slides VP).')
    novos = [copiar(s, prs) for s in parte.slides]
    ids_parte = [sid(s) for s in novos]
    # move a parte para depois do FP13 (ou do último FP), antes do fechamento
    lst = prs.slides._sldIdLst
    ids_xml = list(lst)
    k = max(i for i, x in enumerate(antes) if x.startswith('FP'))
    for el in ids_xml[len(antes):]:
        lst.remove(el)
    ancora = list(lst)[k]
    for el in reversed(ids_xml[len(antes):]):
        ancora.addnext(el)
    roteiro = linha_roteiro(prs, ids_parte)
    # renumera o número impresso no pé (caixa só de dígitos, abaixo de 6,8 pol.)
    renum = 0
    for pos, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and Emu(sh.top).inches > 6.8:
                run = sh.text_frame.paragraphs[0].runs[0]
                if run.text != str(pos):
                    run.text = str(pos)
                    renum += 1
    prs.save(saida)
    depois = [sid(s) for s in Presentation(saida).slides]
    print(f'{len(antes)} slides do relatório + {len(novos)} da parte = {len(depois)} → {saida}')
    print(f'parte entre {antes[k]} e {antes[k + 1] if k + 1 < len(antes) else "o fim"}: {ids_parte[0]}…{ids_parte[-1]}; '
          f'roteiro A03 {"atualizado" if roteiro else "NÃO atualizado"}; {renum} números de página corrigidos')


if __name__ == '__main__':
    main()
