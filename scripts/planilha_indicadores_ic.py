#!/usr/bin/env python
"""Aba "Indicadores" na planilha de gerenciamento de artigos e fichamentos da IC.

    ./.venv/bin/python scripts/planilha_indicadores_ic.py [planilha.xlsx]

Mapeia os 31 indicadores da versão final (00_referencia): definição, origem no Censo, onde são
calculados, números da versão final, onde são usados (tabelas, figuras, versões, slides do deck
da orientadora) e links "Abrir" relativos, como os da aba Artigos.

A aba entra direto no pacote .xlsx (XML da planilha, da tabela e dos links), sem regravar o
resto: o openpyxl não preserva gráficos, e a aba Dashboard tem três. Os estilos são os da aba
Artigos (cabeçalho, texto e link). Rodar de novo substitui a aba, não duplica.
"""
import json
import os
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import quote
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import eda_versoes as ev  # noqa: E402
import verificar_versao as vv  # noqa: E402

RAIZ = ev.RAIZ
A = RAIZ / 'analises_exploratorias'
REF = A / '00_referencia'
PLANILHA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ.parent / 'Relatorios e Andamento' / 'gerenciamento_artigos_cientificos_IC.xlsx'
DECK = RAIZ / 'docs' / 'Apresentacoes_IVS' / 'Apresentacao_Orientadora_EDA_Fatorial_2026-10.pptx'
QUADRO_XLSX = RAIZ / 'banco_de_dados' / 'entrega_orientadora' / 'Quadro_Indicadores.xlsx'
ABA, TABELA = 'Indicadores', 'TabelaIndicadores'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CALCULO_EXTRA = {'renda_media_sem_extremo': ('src/ivs_censo/renda.py', 'renda_sem_extremos'),
                 'renda_media_mediana_mun': ('src/ivs_censo/renda.py', 'renda_imputada_mediana_municipal'),
                 'banheiro_graduado': ('scripts/ivs_especificacoes.py', 'banheiro_graduado')}
FIG_DIM = {'Saneamento': ('F02', 'F03', 'F04', 'F32'), 'Socioeconômica': ('F05', 'F06', 'F07', 'F33'),
           'Morfologia': ('F08', 'F09', 'F10', 'F34'), 'Demografia': ('F11', 'F12', 'F13', 'F35')}


def tab(id_):
    return pd.read_csv(next((REF / 'tabelas').glob(f'{id_}_*.csv')))


def rel(p: Path) -> str:
    """Caminho relativo à planilha, codificado como nos links da aba Artigos."""
    return quote(os.path.relpath(p, PLANILHA.parent))


def linha_codigo(codigo):
    if codigo in CALCULO_EXTRA:
        arq, func = CALCULO_EXTRA[codigo]
        txt = (RAIZ / arq).read_text(encoding='utf-8').splitlines()
        n = next(i for i, l in enumerate(txt, 1) if l.startswith(f'def {func}('))
        return arq, f'{arq}, função {func} (linha {n})'
    txt = (RAIZ / 'src/ivs_censo/indicadores.py').read_text(encoding='utf-8').splitlines()
    n = next(i for i, l in enumerate(txt, 1) if f"Indicador('{codigo}'" in l)
    return 'src/ivs_censo/indicadores.py', f'src/ivs_censo/indicadores.py, linha {n}'


def montar() -> tuple[list, list]:
    cfg = json.loads((REF / 'config.json').read_text(encoding='utf-8'))
    comp = {c if c != 'renda_inv' else 'renda_media' for c in cfg['indice']['colunas']}
    q = pd.read_csv(ev.QUADRO, sep=';', encoding='utf-8-sig').set_index('indicador')
    ext = json.loads(ev.MODELOS.read_text(encoding='utf-8'))['dicionario_extra']
    t02, t03, t15 = tab('T02').set_index('indicador'), tab('T03').set_index('indicador'), tab('T15').set_index('indicador')
    cat = pd.read_csv(REF / 'catalogo.csv', dtype=str).set_index('id')
    tabelas_txt = {f.name.split('_')[0]: f.read_text(encoding='utf-8') for f in sorted((REF / 'tabelas').glob('T*.csv'))}
    texto, notas, _, _ = vv.textos_pptx(DECK)
    versoes = []
    for c in sorted(A.glob('**/config.json')):
        v = json.loads(c.read_text(encoding='utf-8'))
        ic = v.get('indice') or {}
        conj = lambda i: {('renda_media' if i.get('renda', 'original') == 'original' else
                           'renda_media_' + i['renda']) if x == 'renda_inv' else x for x in (i.get('colunas') or [])}
        cols = conj(ic)
        alt = set().union(*[conj({**ic, **(a.get('indice') or {})}) for a in v.get('sensibilidade', [])]) - cols
        if v.get('analise') or c.parent.name == '00_sintese':
            continue
        versoes.append((str(c.parent.relative_to(A)), v.get('indicadores', 'todos'), cols, alt))
    cab = ['ID', 'Indicador (código)', 'Nome', 'Dimensão', 'Entra no IVS-7', 'Componente do índice da referência', 'O que mede',
           'Numerador', 'Denominador', 'Complemento', 'Arquivo do Censo de origem', 'Abrir origem', 'Calculado em', 'Abrir código',
           'n (versão final)', 'Média', 'Mediana', '% zeros', '% faltantes', 'AUC contra FCU', 'Tabelas da versão final',
           'Figuras da versão final', 'Abrir figura', 'Descrito nas versões', 'Componente do índice nas versões (T34 = alternativa testada)',
           'Slides do deck da orientadora', 'Abrir deck', 'Quadro de indicadores', 'Observação']
    linhas = []
    for k, cod in enumerate(t02.index, 1):
        rot, dim = ev.ROT[cod], t02.loc[cod, 'dimensao']
        if cod in q.index:
            d = q.loc[cod]
            mede, num, den, compl, fonte, obs = d['o_que_mede'], d['numerador'], d['denominador'], d['complemento'], d['arquivo_fonte'], d['observacao']
        else:
            e = ext[cod]
            mede, num, den, compl, fonte, obs = e['o_que_mede'], e['numerador_curto'], e['denominador_curto'], 'Não', 'derivado de outra coluna da base', 'variante criada depois do quadro de indicadores'
        arqs = [f.strip() for f in str(fonte).split('|') if (RAIZ / 'dados' / f.strip()).exists()]
        arq_cod, onde = linha_codigo(cod)
        tabs = [t for t, txt in tabelas_txt.items() if re.search(rf'\b{re.escape(cod)}\b', txt) or rot in txt]
        figs = ['F14', 'F16', 'F17', 'F19'] + [f for f in FIG_DIM[dim] if f in cat.index]
        if cod in comp or cod in ev.DEMANDAS:
            figs.append('F31')
        if cod in ev.ENVELHECIMENTO:
            figs += ['F20', 'F21']
        if cod in comp:
            figs += ['F22', 'F23', 'F24', 'F25']
        fig_hist = REF / cat.loc[FIG_DIM[dim][0], 'arquivo']
        desc = [v for v, inds, _, _ in versoes if inds == 'todos' or cod in inds]
        compv = [v for v, _, cols, _ in versoes if cod in cols] + [f'{v} (T34)' for v, _, _, alt in versoes if cod in alt]
        slides = [str(i + 1) for i, t in enumerate(texto) if rot in t]
        num_val = lambda x, d=4: None if pd.isna(x) else round(float(x), d)
        linhas.append([
            f'IND{k:02d}', cod, rot, dim, 'Sim' if (cod in q.index and q.loc[cod, 'entra_no_ivs7'] == 'Sim') else 'Não',
            'Sim' if cod in comp else 'Não', mede, num, den, compl, fonte,
            ('link', 'Abrir arquivo', RAIZ / 'dados' / arqs[0]) if arqs else '—',
            onde, ('link', 'Abrir código', RAIZ / arq_cod),
            num_val(t03.loc[cod, 'n'], 0), num_val(t03.loc[cod, 'media'], 2 if cod.startswith('renda') else 4),
            num_val(t03.loc[cod, 'mediana'], 2 if cod.startswith('renda') else 4), num_val(t03.loc[cod, 'pct_zeros'], 1),
            num_val(t02.loc[cod, 'pct_faltantes'], 1), num_val(t15.loc[cod, 'auc']) if cod in t15.index else None,
            ', '.join(tabs), ', '.join(sorted(set(figs))), ('link', 'Abrir histograma', fig_hist),
            f'{len(desc)} versões: ' + ', '.join(desc), (f'{len(compv)} versões: ' + ', '.join(compv)) if compv else 'nenhuma',
            ', '.join(slides) if slides else '—', ('link', 'Abrir deck', DECK), ('link', 'Abrir quadro', QUADRO_XLSX),
            '' if pd.isna(obs) else obs])
    return cab, linhas


# ── Escrita direta no pacote .xlsx ────────────────────────────────────────────
def estilos_da_aba_artigos(z):
    wb = z.read('xl/workbook.xml').decode('utf-8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rid = re.search(r'<sheet [^>]*name="Artigos"[^>]*r:id="(rId\d+)"', wb).group(1)
    alvo = re.search(rf'<Relationship [^>]*Id="{rid}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{rid}"', rels)
    xml = z.read('xl/' + alvo.group(1).lstrip('/').replace('xl/', '')).decode('utf-8')
    s = lambda ref: (re.search(rf'<c r="{ref}"[^>]*?\bs="(\d+)"', xml) or [None, '0'])[1]
    return s('A1'), s('B2'), s('N2'), s('D2')


def xml_da_aba(cab, linhas, est_cab, est_txt, est_link, est_num):
    n_col = len(cab)
    larg = [9, 24, 24, 15, 10, 12, 42, 48, 30, 11, 34, 14, 34, 14, 11, 11, 11, 9, 10, 10, 34, 30, 15, 40, 40, 24, 12, 14, 36]
    cols = ''.join(f'<col min="{i}" max="{i}" width="{w}" customWidth="1"/>' for i, w in enumerate(larg[:n_col], 1))
    linhas_xml, links = [], []
    cel = lambda ref, v, s: (f'<c r="{ref}" s="{s}" t="inlineStr"><is><t xml:space="preserve">{escape(str(v))}</t></is></c>'
                             if isinstance(v, str) else (f'<c r="{ref}" s="{est_num}"><v>{v}</v></c>' if v is not None else f'<c r="{ref}" s="{s}"/>'))
    linhas_xml.append('<row r="1">' + ''.join(cel(f'{get_column_letter(j)}1', c, est_cab) for j, c in enumerate(cab, 1)) + '</row>')
    for i, ln in enumerate(linhas, 2):
        cs = []
        for j, v in enumerate(ln, 1):
            ref = f'{get_column_letter(j)}{i}'
            if isinstance(v, tuple):
                links.append((ref, rel(v[2])))
                cs.append(cel(ref, v[1], est_link))
            else:
                cs.append(cel(ref, v, est_txt))
        linhas_xml.append(f'<row r="{i}">' + ''.join(cs) + '</row>')
    ult = f'{get_column_letter(n_col)}{len(linhas) + 1}'
    hl = ''.join(f'<hyperlink ref="{r}" r:id="rIdL{k}"/>' for k, (r, _) in enumerate(links, 1))
    sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="{NS_R}">'
             f'<dimension ref="A1:{ult}"/>'
             '<sheetViews><sheetView workbookViewId="0"><pane xSplit="3" ySplit="1" topLeftCell="D2" activePane="bottomRight" state="frozen"/>'
             '<selection pane="bottomRight" activeCell="D2" sqref="D2"/></sheetView></sheetViews>'
             '<sheetFormatPr defaultRowHeight="15"/>'
             f'<cols>{cols}</cols><sheetData>{"".join(linhas_xml)}</sheetData>'
             f'<hyperlinks>{hl}</hyperlinks>'
             '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
             '<tableParts count="1"><tablePart r:id="rIdT1"/></tableParts></worksheet>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + ''.join(f'<Relationship Id="rIdL{k}" Type="{NS_R}/hyperlink" Target="{escape(t)}" TargetMode="External"/>'
                      for k, (_, t) in enumerate(links, 1))
            + f'<Relationship Id="rIdT1" Type="{NS_R}/table" Target="../tables/TABELA_XML"/></Relationships>')
    return sheet, rels, f'A1:{ult}', len(links)


def xml_da_tabela(tid, ref, cab):
    cols = ''.join(f'<tableColumn id="{i}" name="{escape(c)}"/>' for i, c in enumerate(cab, 1))
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" id="{tid}" name="{TABELA}" '
            f'displayName="{TABELA}" ref="{ref}" totalsRowShown="0"><autoFilter ref="{ref}"/>'
            f'<tableColumns count="{len(cab)}">{cols}</tableColumns>'
            '<tableStyleInfo name="TableStyleMedium2" showFirstColumn="0" showLastColumn="0" showRowStripes="1" showColumnStripes="0"/></table>')


def gravar(cab, linhas):
    with zipfile.ZipFile(PLANILHA) as z:
        itens = {n: z.read(n) for n in z.namelist()}
        est = estilos_da_aba_artigos(z)
    wb = itens['xl/workbook.xml'].decode('utf-8')
    wrels = itens['xl/_rels/workbook.xml.rels'].decode('utf-8')
    tipos = itens['[Content_Types].xml'].decode('utf-8')
    # rodar de novo substitui a aba anterior
    m = re.search(rf'<sheet [^>]*name="{ABA}"[^>]*r:id="(rId\d+)"\s*/>', wb)
    if m:
        alvo = re.search(rf'<Relationship [^>]*Id="{m.group(1)}"[^>]*Target="([^"]+)"[^>]*/>', wrels).group(1)
        nome = 'xl/' + alvo.replace('/xl/', '').lstrip('/')
        rels_antigo = nome.replace('worksheets/', 'worksheets/_rels/') + '.rels'
        tabela_antiga = re.search(r'Target="\.\./tables/([^"]+)"', itens[rels_antigo].decode('utf-8')).group(1)
        for n in (nome, rels_antigo, f'xl/tables/{tabela_antiga}'):
            itens.pop(n, None)
            tipos = re.sub(rf'<Override PartName="/{re.escape(n)}"[^>]*/>', '', tipos)
        wb = wb.replace(m.group(0), '')
        wrels = re.sub(rf'<Relationship [^>]*Id="{m.group(1)}"[^>]*/>', '', wrels)
    n_sheet = max(int(x) for x in re.findall(r'xl/worksheets/sheet(\d+)\.xml', ' '.join(itens))) + 1
    n_tab = max([int(x) for x in re.findall(r'xl/tables/table(\d+)\.xml', ' '.join(itens))] + [0]) + 1
    ids_tab = [int(x) for n, b in itens.items() if n.startswith('xl/tables/') for x in re.findall(rb'<table [^>]*\bid="(\d+)"', b)]
    sheet, srels, ref, n_links = xml_da_aba(cab, linhas, *est)
    itens[f'xl/worksheets/sheet{n_sheet}.xml'] = sheet.encode('utf-8')
    itens[f'xl/worksheets/_rels/sheet{n_sheet}.xml.rels'] = srels.replace('TABELA_XML', f'table{n_tab}.xml').encode('utf-8')
    itens[f'xl/tables/table{n_tab}.xml'] = xml_da_tabela(max(ids_tab + [0]) + 1, ref, cab).encode('utf-8')
    rid = 'rId' + str(max(int(x) for x in re.findall(r'Id="rId(\d+)"', wrels)) + 1)
    wrels = wrels.replace('</Relationships>', f'<Relationship Id="{rid}" Type="{NS_R}/worksheet" Target="worksheets/sheet{n_sheet}.xml"/></Relationships>')
    sid = max(int(x) for x in re.findall(r'<sheet [^>]*sheetId="(\d+)"', wb)) + 1
    wb = wb.replace('</sheets>', f'<sheet name="{ABA}" sheetId="{sid}" r:id="{rid}"/></sheets>')
    tipos = tipos.replace('</Types>',
                          f'<Override PartName="/xl/worksheets/sheet{n_sheet}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                          f'<Override PartName="/xl/tables/table{n_tab}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"/></Types>')
    itens['xl/workbook.xml'], itens['xl/_rels/workbook.xml.rels'], itens['[Content_Types].xml'] = wb.encode('utf-8'), wrels.encode('utf-8'), tipos.encode('utf-8')
    tmp = PLANILHA.with_suffix('.tmp.xlsx')
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in ['[Content_Types].xml'] + [n for n in itens if n != '[Content_Types].xml']:
            z.writestr(n, itens[n])
    os.replace(tmp, PLANILHA)
    return n_links


def main():
    cab, linhas = montar()
    n_links = gravar(cab, linhas)
    print(f'{PLANILHA.name}: aba {ABA} com {len(linhas)} indicadores, {len(cab)} colunas e {n_links} links')


if __name__ == '__main__':
    main()
