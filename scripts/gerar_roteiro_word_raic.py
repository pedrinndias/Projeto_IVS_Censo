"""Roteiro em Word da apresentação oral da RAIC 2026, para o ensaio com a orientadora.

Lê as falas, as telas, os tempos e as perguntas de `relatorio_parcial_raic/ROTEIRO_APRESENTACAO_ORAL.md`
(a mesma fonte das notas do pptx) e usa como miniaturas o `_ENVIO.pptx` como está — inclusive com as
edições feitas à mão. O design segue o template da Coordic: Calibri, a faixa da 34ª RAIC e as cores do
cabeçalho.

Marcação das falas no roteiro: **âncora** (negrito), *(pausa)*, **(decorar)** e ==trecho== (depende
da orientadora; sai com marca-texto amarelo).

Uso:  ./.venv/bin/python scripts/gerar_roteiro_word_raic.py
Requer o LibreOffice (miniaturas) e o `sips` do macOS.
Saída: relatorio_parcial_raic/Roteiro_Apresentacao_Oral_RAIC_2026_Pedro.docx
"""
from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader, PdfWriter

RAIZ = Path(__file__).resolve().parents[1]
PASTA = RAIZ / 'relatorio_parcial_raic'
ROTEIRO = PASTA / 'ROTEIRO_APRESENTACAO_ORAL.md'
DECK = PASTA / 'Apresentacao_Oral_RAIC_2026_Pedro_ENVIO.pptx'
MINIATURAS = PASTA / 'figuras' / 'miniaturas'
SAIDA = PASTA / 'Roteiro_Apresentacao_Oral_RAIC_2026_Pedro.docx'
SOFFICE = '/Applications/LibreOffice.app/Contents/MacOS/soffice'

# Paleta do template (a mesma de scripts/gerar_apresentacao_oral_raic.py).
TEAL, TEAL_CLARO, CORAL, LARANJA = '3B848E', '79BCC4', 'D8634E', 'E8813A'
LAVANDA, FUNDO, TEXTO, CINZA = 'E1D6E9', 'E9F4F4', '1F2A33', '6B7780'
PAUSA = '7A5C99'

# Critério de avaliação do edital que cada slide atende (edital RAIC 2026, "Critérios de Avaliação").
CRITERIO = {1: 'Conteúdo obrigatório: identificação', 2: 'Clareza e qualidade da apresentação',
            3: 'Conteúdo obrigatório: objetivos', 4: 'Metodologia · domínio do tema',
            5: 'Metodologia · domínio do tema', 6: 'Discussão dos resultados', 7: 'Discussão dos resultados',
            8: 'Discussão dos resultados', 9: 'Discussão dos resultados', 10: 'Conclusões',
            11: 'Passagem para a arguição'}


# ---------------------------------------------------------------- leitura do roteiro

def ler_roteiro() -> tuple[list[dict], list[dict], list[dict]]:
    md = ROTEIRO.read_text(encoding='utf-8')
    corpo, resto = md.split('## Reserva para a arguição', 1)
    slides = []
    for bloco in re.split(r'(?m)^### ', corpo.split('## Os slides', 1)[1])[1:]:
        cab, texto = bloco.split('\n', 1)
        m = re.match(r'(\d+)\. (?:([A-ZÇÕÃÉ]+) · )?(.*?) — (\d+) s', cab)
        campos = {}
        for rot, val in re.findall(r'(?ms)^- \*\*(Tela|Visual|Fala[^*]*):\*\* (.*?)(?=^- \*\*|\Z)', texto):
            val = val.split('\n---', 1)[0]  # o último slide encosta no separador da seção seguinte
            campos['fala' if rot.startswith('Fala') else 'tela'] = re.sub(r'\s+', ' ', val).strip()
        slides.append({'n': int(m[1]), 'secao': m[2], 'titulo': m[3], 'seg': int(m[4]), **campos})
    tabela, perguntas_md = resto.split('## Perguntas prováveis', 1)
    reservas = [{'id': c[0], 'slide': c[1].replace('**', ''), 'pergunta': c[2].strip('"')}
                for c in (re.findall(r'\| (R\d) \| (.*?) \| (.*?) \|', tabela))]
    perguntas_md = perguntas_md.split('\n---', 1)[0]
    perguntas = []
    for q, a in re.findall(r'(?ms)^\d+\. \*\*"(.*?)"\*\* (.*?)(?=^\d+\. |\Z)', perguntas_md):
        a = re.sub(r'\s+', ' ', a).strip()
        r = re.findall(r'\((R\d)\)', a)
        perguntas.append({'p': q, 'r': re.sub(r'\s*\(R\d\)', '', a), 'reserva': ', '.join(r) or '—'})
    return slides, reservas, perguntas


# ---------------------------------------------------------------- miniaturas

def miniaturas() -> list[Path]:
    """Uma PNG por slide do _ENVIO.pptx (LibreOffice → PDF → página a página → sips)."""
    MINIATURAS.mkdir(parents=True, exist_ok=True)
    for velha in MINIATURAS.glob('slide_*.png'):
        velha.unlink()
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([SOFFICE, '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(DECK)],
                       check=True, capture_output=True)
        leitor = PdfReader(next(Path(tmp).glob('*.pdf')))
        caminhos = []
        for i, pagina in enumerate(leitor.pages, 1):
            um = Path(tmp) / f'p{i:02d}.pdf'
            w = PdfWriter()
            w.add_page(pagina)
            w.write(str(um))
            png = MINIATURAS / f'slide_{i:02d}.png'
            subprocess.run(['sips', '-s', 'format', 'png', '--resampleWidth', '2000', str(um), '--out', str(png)],
                           check=True, capture_output=True)
            caminhos.append(png)
    return caminhos


def faixa(png: Path, nome: str, fracao: float) -> Path:
    """Recorta a faixa de cima de um slide (o cabeçalho da 34ª RAIC) para usar no Word."""
    im = Image.open(png)
    saida = MINIATURAS / nome
    im.crop((0, 0, im.width, int(im.height * fracao))).save(saida)
    return saida


# ---------------------------------------------------------------- utilidades do docx

def sombrear(elemento, cor: str) -> None:
    pr = elemento._tc.get_or_add_tcPr() if hasattr(elemento, '_tc') else elemento._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor)
    pr.append(shd)


def bordas_celula(celula, cor: str = CINZA, tam: int = 6, lados=('top', 'left', 'bottom', 'right')) -> None:
    tcpr = celula._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single' if lado in lados else 'nil')
        e.set(qn('w:sz'), str(tam))
        e.set(qn('w:color'), cor)
        b.append(e)
    tcpr.append(b)


def sem_bordas(tabela) -> None:
    for linha in tabela.rows:
        for c in linha.cells:
            bordas_celula(c, lados=())


def larguras(tabela, cms: list[float]) -> None:
    """Larguras fixas na grade e em cada célula (o LibreOffice ignora só a da célula)."""
    tabela.autofit = False
    for i, w in enumerate(cms):
        tabela.columns[i].width = Cm(w)
        for c in tabela.columns[i].cells:
            c.width = Cm(w)


def run(p, texto: str, tam: float | None = None, bold=False, cor: str | None = TEXTO, italic=False,
        destaque=False):
    r = p.add_run(texto)
    r.bold, r.italic = bold, italic
    if tam:
        r.font.size = Pt(tam)
    if cor:
        r.font.color.rgb = RGBColor.from_string(cor)
    if destaque:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def par(container, texto: str = '', tam: float = 11, bold=False, cor: str = TEXTO, alinhar=None, antes=0,
        depois=4, italic=False, entre=None):
    p = container.add_paragraph()
    if texto:
        run(p, texto, tam, bold, cor, italic=italic)
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(antes), Pt(depois)
    if entre:
        f.line_spacing = entre
    if alinhar is not None:
        p.alignment = alinhar
    return p


def markdown(p, texto: str, tam: float, cor: str = TEXTO) -> None:
    """Escreve no parágrafo o texto do roteiro com **negrito**, *(pausa)*, **(decorar)** e ==destaque==."""
    for peca in re.split(r'(\*\*\(decorar\)\*\*|\*\([^)]*\)\*|==.*?==|\*\*.*?\*\*)', texto):
        if not peca:
            continue
        if peca == '**(decorar)**':
            run(p, ' DECORAR ▸ ', tam - 3, bold=True, cor=CORAL)
        elif peca.startswith('*(') and peca.endswith(')*'):
            run(p, f'({peca[2:-2]})', tam - 1, italic=True, cor=PAUSA)
        elif peca.startswith('=='):
            markdown_destaque(p, peca[2:-2], tam, cor)
        elif peca.startswith('**'):
            run(p, peca[2:-2], tam, bold=True, cor=cor)
        else:
            run(p, peca, tam, cor=cor)


def markdown_destaque(p, texto: str, tam: float, cor: str) -> None:
    for peca in re.split(r'(\*\*.*?\*\*)', texto):
        if peca:
            run(p, peca.strip('*'), tam, bold=peca.startswith('**'), cor=cor, destaque=True)


def caixa(container, cor_fundo: str | None, cor_borda: str | None = None, largura: float = 17.0):
    t = container.add_table(rows=1, cols=1)
    larguras(t, [largura])
    c = t.cell(0, 0)
    if cor_fundo:
        sombrear(c, cor_fundo)
    bordas_celula(c, cor_borda or cor_fundo or 'FFFFFF', 8)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    return c


def quebra(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def campo_pagina(p) -> None:
    for tipo, texto in (('begin', None), (None, 'PAGE'), ('end', None)):
        r = p.add_run()
        r.font.size, r.font.color.rgb = Pt(9), RGBColor.from_string(CINZA)
        if tipo:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tipo)
        else:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = texto
        r._r.append(e)


# ---------------------------------------------------------------- páginas

def preparar(doc) -> None:
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin, s.bottom_margin = Cm(1.6), Cm(1.6)
    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    estilo.font.size = Pt(11)
    rodape = s.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(rodape, 'Roteiro da apresentação oral · RAIC 2026 · Pedro Dias Soares · ', 9, cor=CINZA)
    campo_pagina(rodape)


def capa(doc, banner: Path) -> None:
    doc.add_picture(str(banner), width=Cm(17))
    par(doc, 'ROTEIRO DA APRESENTAÇÃO ORAL', 13, True, TEAL, antes=28, depois=2)
    par(doc, 'Ensaio com a orientadora', 30, True, TEXTO, depois=14)
    p = par(doc, depois=24, entre=1.15)
    run(p, 'Índice de vulnerabilidade à saúde das cidades brasileiras: dados do censo demográfico de 2022',
        16, cor=TEXTO)
    c = caixa(doc, FUNDO)
    for rot, val in [('Aluno', 'Pedro Dias Soares'), ('Orientadora', 'Profa. Amanda Cristina de Souza Andrade'),
                     ('Grupo de pesquisa', 'NESPE'), ('Programa', 'FAPEMIG · ingresso em 06/03/2026'),
                     ('Evento', '34ª RAIC — Fiocruz Minas, 9 a 13/11/2026 · modalidade oral')]:
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run(p, f'{rot}:  ', 12, True, TEAL)
        run(p, val, 12)
    par(doc, 'Documento interno de ensaio — não vai para a Coordic.', 10, cor=CINZA, antes=18, italic=True)
    par(doc, 'Belo Horizonte, outubro de 2026', 11, True, TEXTO, antes=40)


def para_professora(doc, slides) -> None:
    quebra(doc)
    par(doc, 'PARA A PROFESSORA AMANDA', 12, True, TEAL, depois=6)
    p = par(doc, depois=10, entre=1.25)
    run(p, 'Professora Amanda, este é o roteiro da fala que vou ensaiar. Quero saber, principalmente, ', 12)
    run(p, 'onde eu me perco', 12, True)
    run(p, ' e se ', 12)
    run(p, 'alguma frase diz mais do que os dados sustentam', 12, True)
    run(p, '. Proponho que a fala seja ouvida duas vezes, como abaixo. No fim da página deixei uma proposta '
           'que depende da sua leitura.', 12)

    t = doc.add_table(rows=1, cols=2)
    larguras(t, [8.4, 8.4])
    sem_bordas(t)
    for c, (n, quem, perg) in zip(t.rows[0].cells, [
            ('1ª passada', 'como avaliadora de fora da área', 'Entendi? Onde me perdi? Algum termo ficou sem '
                                                              'explicação?'),
            ('2ª passada', 'como especialista', 'O rigor se sustenta? Alguma frase afirma mais do que os dados '
                                                'mostram?')]):
        sombrear(c, LAVANDA)
        p = c.paragraphs[0]
        run(p, n.upper(), 10, True, TEAL)
        p2 = c.add_paragraph()
        run(p2, quem, 13, True)
        p3 = c.add_paragraph()
        p3.paragraph_format.space_after = Pt(4)
        run(p3, perg, 11)
    par(doc, 'O que já decidi (para confirmar)', 13, True, TEAL, antes=14, depois=4)
    t = doc.add_table(rows=0, cols=2)
    for item, decisao in [('Grupo de pesquisa na capa', 'NESPE'),
                          ('Objetivo geral do projeto (slide 3)', 'o texto do relatório parcial: "Avaliar as '
                           'desigualdades intraurbanas por meio da construção de um Índice de Vulnerabilidade à '
                           'Saúde (IVS) de cidades brasileiras."'),
                          ('Linha sobre a análise fatorial no slide 10', 'retirada da tela; a fala cita a análise '
                                                                        'como em andamento'),
                          ('Reserva R7 (análise fatorial)', 'retirada do arquivo; se a banca perguntar, respondo '
                                                            'com a frase do slide 10')]:
        c1, c2 = t.add_row().cells
        run(c1.paragraphs[0], item, 11, True)
        run(c2.paragraphs[0], decisao, 11)
        for c in (c1, c2):
            bordas_celula(c, TEAL_CLARO, 4, ('bottom',))
    larguras(t, [5.6, 11.2])
    par(doc, 'Uma proposta para a sua leitura', 13, True, TEAL, antes=14, depois=4)
    for titulo, proposta, pergunta in [
            ('A frase sobre a análise fatorial (fala do slide 10)',
             'A análise fatorial começou entre agosto e setembro, e os resultados ainda são preliminares, em '
             'discussão com a professora Amanda.',
             'É o texto do relatório parcial. Serve para a banca, ou prefere outra formulação?')]:
        c = caixa(doc, None, CORAL)
        p = c.paragraphs[0]
        run(p, titulo, 12, True, CORAL)
        p2 = c.add_paragraph()
        run(p2, 'Minha proposta: ', 11, True)
        run(p2, proposta, 11, destaque=True)
        p3 = c.add_paragraph()
        p3.paragraph_format.space_after = Pt(4)
        run(p3, pergunta, 11, cor=CINZA)
        par(doc, depois=2)
    par(doc, 'Critérios da RAIC: clareza e qualidade da apresentação · domínio do tema · discussão dos '
             'resultados · capacidade de responder às perguntas da banca.', 10, cor=CINZA, antes=6)


def fio_e_tempo(doc, slides) -> None:
    quebra(doc)
    par(doc, 'O FIO DA APRESENTAÇÃO', 12, True, TEAL, depois=6)
    c = caixa(doc, None, TEAL)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run(p, 'A média da cidade esconde o que acontece dentro dela. Hoje mostro a etapa que vem antes do índice: '
           'o que o Censo 2022 me obrigou a decidir. Os dados me mostraram três coisas, e cada uma muda o '
           'cálculo.', 14, True, TEAL)
    total = sum(s['seg'] for s in slides)
    par(doc, 'O tempo', 13, True, TEAL, antes=14, depois=4)
    p = par(doc, depois=6, entre=1.2)
    run(p, 'Limite da RAIC: 10 min de fala e cerca de 10 min de perguntas. ', 11)
    run(p, f'Estimativa: {total // 60} min {total % 60:02d} s', 11, True)
    run(p, ' a 130 palavras por minuto, com as pausas. Alvo: de 7 a 9 min. A estimativa é margem para o '
           'nervosismo, não meta: se o ensaio passar de 10 min, corto conteúdo em vez de acelerar.', 11)
    t = doc.add_table(rows=1, cols=4)
    for c, h in zip(t.rows[0].cells, ['Slide', 'Título', 'Tempo', 'Critério da RAIC']):
        sombrear(c, TEAL)
        run(c.paragraphs[0], h, 10, True, 'FFFFFF')
    for s in slides:
        cs = t.add_row().cells
        for c, v in zip(cs, [str(s['n']), s['titulo'], f"{s['seg']} s", CRITERIO[s['n']]]):
            run(c.paragraphs[0], v, 10)
            bordas_celula(c, LAVANDA, 4, ('bottom',))
    larguras(t, [1.3, 8.6, 1.6, 5.3])
    par(doc, 'Como ler as falas', 13, True, TEAL, antes=14, depois=4)
    p = par(doc, depois=3)
    run(p, 'negrito', 11, True)
    run(p, ' = âncora: o que eu olho no papel para lembrar a frase.   ', 11)
    run(p, '(pausa)', 11, italic=True, cor=PAUSA)
    run(p, ' = parar um segundo.', 11)
    p = par(doc, depois=3)
    run(p, 'DECORAR ▸', 9, True, CORAL)
    run(p, ' = sai de cor: só a abertura (slides 1 e 2) e o fecho (fim do 10). O resto eu falo a partir do '
           'título do slide.   ', 11)
    p = par(doc, depois=3)
    run(p, 'amarelo', 11, destaque=True)
    run(p, ' = depende da sua leitura.', 11)


def pagina_slide(doc, s, png: Path) -> None:
    quebra(doc)
    t = doc.add_table(rows=1, cols=2)
    larguras(t, [10.6, 6.2])
    esq, dir_ = t.rows[0].cells
    for c in (esq, dir_):
        sombrear(c, TEAL)
        bordas_celula(c, TEAL, 4)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    secao = f' · {s["secao"]}' if s['secao'] else ''
    run(esq.paragraphs[0], f'SLIDE {s["n"]}{secao}  ·  {s["seg"]} s', 12, True, 'FFFFFF')
    dir_.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(dir_.paragraphs[0], CRITERIO[s['n']], 9, cor='FFFFFF')
    par(doc, s['titulo'], 16, True, TEXTO, antes=8, depois=6)

    t = doc.add_table(rows=1, cols=2)
    larguras(t, [12.4, 4.4])
    sem_bordas(t)
    corpo, margem = t.rows[0].cells
    foto = corpo.paragraphs[0]
    foto.add_run().add_picture(str(png), width=Cm(11.9))
    foto.paragraph_format.space_after = Pt(2)
    if s.get('tela'):
        p = corpo.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run(p, 'NA TELA  ', 8, True, CINZA)
        markdown(p, s['tela'], 9, CINZA)
    p = corpo.add_paragraph()
    run(p, 'FALA', 9, True, TEAL)
    p.paragraph_format.space_after = Pt(2)
    fala = s.get('fala', '').strip().strip('"')
    p = corpo.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    markdown(p, fala, 13)

    sombrear(margem, 'FAFAFA')
    bordas_celula(margem, LAVANDA, 8, ('left',))
    p = margem.paragraphs[0]
    run(p, 'OBSERVAÇÕES', 9, True, TEAL)
    for _ in range(16):
        q = margem.add_paragraph()
        q.paragraph_format.space_after = Pt(9)
        run(q, '_' * 22, 9, cor='C9CED2')
    q = margem.add_paragraph()
    run(q, 'Tempo real: ______', 10, True, CINZA)


def arguicao(doc, perguntas) -> None:
    quebra(doc)
    par(doc, 'ARGUIÇÃO — ENSAIO DAS PERGUNTAS', 12, True, TEAL, depois=4)
    par(doc, 'Peço que a professora faça as perguntas sem aviso, na ordem que quiser. A primeira é a mais '
             'provável: "cadê o índice?". A resposta é curta e, quando houver, aponta o slide de reserva.', 11,
        depois=8, entre=1.2)
    t = doc.add_table(rows=1, cols=3)
    for c, h in zip(t.rows[0].cells, ['Pergunta', 'Minha resposta curta', 'Reserva']):
        sombrear(c, TEAL)
        run(c.paragraphs[0], h, 10, True, 'FFFFFF')
    for k, q in enumerate(perguntas):
        cs = t.add_row().cells
        run(cs[0].paragraphs[0], q['p'], 10.5, True)
        run(cs[1].paragraphs[0], q['r'], 10.5)
        run(cs[2].paragraphs[0], q['reserva'], 10.5, True, TEAL)
        for c in cs:
            if k % 2:
                sombrear(c, FUNDO)
            bordas_celula(c, LAVANDA, 4, ('bottom',))
    larguras(t, [5.0, 10.2, 1.6])


def reservas(doc, lista, pngs) -> None:
    quebra(doc)
    par(doc, 'SLIDES DE RESERVA (DEPOIS DO "OBRIGADO")', 12, True, TEAL, depois=4)
    par(doc, 'Só aparecem se a banca perguntar.', 11, depois=8)
    t = doc.add_table(rows=0, cols=2)
    sem_bordas(t)
    for i in range(0, len(lista), 2):
        cs = t.add_row().cells
        for c, r in zip(cs, lista[i:i + 2]):
            n = 11 + int(r['id'][1:])
            c.paragraphs[0].add_run().add_picture(str(pngs[n - 1]), width=Cm(7.4))
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run(p, f"{r['id']} · {r['slide']}", 9.5, True, TEAL)
            run(p, f"\n{r['pergunta']}", 9.5, italic=True, cor=CINZA)
        for c in cs:
            bordas_celula(c, lados=())
    larguras(t, [8.4, 8.4])


def main() -> None:
    slides, lista_reservas, perguntas = ler_roteiro()
    pngs = miniaturas()
    banner = faixa(pngs[0], 'faixa_capa.png', 0.27)
    doc = Document()
    preparar(doc)
    capa(doc, banner)
    para_professora(doc, slides)
    fio_e_tempo(doc, slides)
    for s in slides:
        pagina_slide(doc, s, pngs[s['n'] - 1])
    arguicao(doc, perguntas)
    reservas(doc, lista_reservas, pngs)
    doc.save(str(SAIDA))
    print(f'{SAIDA.relative_to(RAIZ)}: {len(slides)} slides de fala, {len(perguntas)} perguntas, '
          f'{len(lista_reservas)} reservas')


if __name__ == '__main__':
    main()
