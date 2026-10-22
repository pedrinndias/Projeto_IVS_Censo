"""Conteúdo da v4 do relatório descritivo: o apoio do Pedro para a reunião (prompt_apoio_reuniao_v4).

Para cada slide da v3: "o que é", 2 a 4 perguntas prováveis com resposta e de onde vem, e o que
saber de cor. Mais P00 (perguntas gerais) e G (glossário). Grava `relatorio_descritivo/apoio/apoio.json`.

Os slides que se repetem (DS, DS-2, FC, eixos VP, DC) saem de modelos preenchidos com os números
lidos do próprio slide da v3. Os de método são escritos à mão; todo número que eles citam tem de
estar no slide (ou nos slides de `ver`), o que `verificar_apoio_relatorio.py` confere.

    uv run --with python-pptx python scripts/dados_apoio_relatorio.py
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation

RAIZ = Path(__file__).resolve().parents[1]
V3 = RAIZ / 'relatorio_descritivo' / 'Relatorio_Descritivo_IVS_2026-10_v3.pptx'
SAIDA = RAIZ / 'relatorio_descritivo' / 'apoio' / 'apoio.json'
ID = re.compile(r'^(A|DM|DS|DC|FA|FC|FP|Z|VP)\d{2}(-\d+)?$')


# ---------------------------------------------------------------- leitura da v3
def ler_v3():
    d = {}
    for s in Presentation(V3).slides:
        sid, textos, linhas = '?', [], []
        for sh in sorted(s.shapes, key=lambda x: (x.top, x.left)):
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if ID.match(t):
                    sid = t
                elif t and not t.isdigit():
                    textos.append(t)
            if sh.has_table:
                linhas += [[c.text.strip() for c in r.cells] for r in sh.table.rows]
        d[sid] = dict(textos=textos, linhas=linhas)
    return d


V = ler_v3()


def cel(sid, rotulo, col=1):
    for ln in V[sid]['linhas']:
        if ln and ln[0] == rotulo:
            return ln[col]
    raise KeyError(f'{sid}: linha "{rotulo}" não encontrada')


def frase(sid, inicio='O que mostra.'):
    for t in V[sid]['textos']:
        if t.startswith(inicio):
            return t[len(inicio):].strip()
    return ''


def titulo(sid):
    t = V[sid]['textos']
    return t[1] if t and t[0].isupper() and len(t) > 1 else t[0]


def num(s):
    return float(re.sub(r'[^\d,-]', '', s.split('(')[0]).replace(',', '.') or 'nan')


def Q(p, r, de):
    return dict(p=p, r=r, de_onde=de)


# ---------------------------------------------------------------- DS: um indicador por slide
POR_QUE_FORA = {
    'Sem água canalizada': 'Foi pedido na demanda 7. Testei nos cenários S3 e S6 e no eixo água das 108 versões; '
                           'não está na referência proposta.',
    'Água não chega': 'Demandas 5 e 7. Testei nos cenários S2 e S6 e no eixo água das 108 versões; não está na '
                      'referência proposta.',
    'Água só no terreno': 'Descreve a canalização junto com "não chega" e "sem canalização". Não testei como '
                          'componente da fatorial.',
    'Sem banheiro': 'Demanda 6. É a V00495, testada no cenário S4 e no eixo banheiro das 108 versões. Na '
                    'referência entrou a forma graduada do banheiro.',
    'Sem banheiro nem sanitário': 'É uma das três partes do sem banheiro. Entra no índice só dentro do banheiro '
                                  'graduado (peso 3), não sozinha.',
    'Só sanitário': 'É uma das três partes do sem banheiro. Entra no índice só dentro do banheiro graduado '
                    '(peso 2), não sozinha.',
    'Banheiro de uso comum': 'É uma das três partes do sem banheiro. Entra no índice só dentro do banheiro '
                             'graduado (peso 1), não sozinha.',
    'Responsável mulher': 'Descreve a população. Não está no IVS-BH 2012 e não testei na fatorial.',
    'Renda sem o extremo': 'É uma variante da renda (demanda 1). Testei nos cenários S0 e S6 e no eixo renda '
                           'das 108 versões; a referência usa a renda original.',
    'Renda (mediana mun.)': 'É uma variante da renda (demanda 1). Testei nos cenários S0 e S6 e no eixo renda '
                            'das 108 versões; a referência usa a renda original.',
    'Moradia não convencional': 'Demanda 9. Testei nos cenários S5 e S6 e no eixo moradia das 108 versões. Tem '
                                'variância baixa (DM10) e não está na referência.',
}
MORADIA = {'Domicílio improvisado', 'Habitação precária', 'Moradia convencional', 'Apartamento', 'Casa',
           'Casa de vila/condomínio'}


def ds(sid):
    nome, sub = titulo(sid), V[sid]['textos'][2]
    dim = V[sid]['textos'][0].capitalize()
    numer, den, arq = cel(sid, 'Numerador'), cel(sid, 'Denominador'), cel(sid, 'Arquivo do Censo')
    unid, ivs7, ref = cel(sid, 'Unidade'), cel(sid, 'Entra no IVS-7?'), cel(sid, 'Entra na referência proposta?')
    n, falt, media, med = cel(sid, 'Setores com valor'), cel(sid, 'Faltantes'), cel(sid, 'Média'), cel(sid, 'Mediana')
    zeros, assim, cv = cel(sid, 'Setores com zero'), cel(sid, 'Assimetria'), cel(sid, 'Coeficiente de variação')
    codigos = re.findall(r'\b[Vv]\d{4,5}\b', numer)
    lista = ', '.join(codigos[:4]) + (f' e mais {len(codigos) - 4}' if len(codigos) > 4 else '')
    dcod = (re.findall(r'\b[Vv]\d{4,5}\b', den) or [den])[0]
    o = (f'{sub}. Cada setor recebe um valor ({unid}); o slide mostra como esse valor se distribui nos '
         f'setores urbanos com dados divulgados. Dimensão: {dim.lower()}.')
    origem = (f'Do Censo 2022, agregados por setor ({arq.split(" | ")[0]}).' if '.csv' in arq
              else 'É calculado a partir de outras colunas da base, que vêm do Censo 2022.')
    dtxt = 'não tem denominador (é o valor médio do setor)' if 'sem denominador' in den else f'denominador: {dcod}'
    q = [Q('De onde vem esse indicador?', f'{origem} Numerador: {lista or numer}; {dtxt}.',
           f'{sid}, tabela "Como é calculado"; Quadro_Indicadores.csv; aba Indicadores da planilha de gerenciamento')]
    z = num(zeros)
    if z >= 50:
        q.append(Q('Por que a mediana é zero?',
                   f'Em {zeros} dos setores nenhum domicílio está nessa situação: a condição se concentra em poucos '
                   f'setores. Por isso descrevo também a média ({media}) e o % de setores com valor acima de zero '
                   f'(slide seguinte).', f'{sid}, tabela de distribuição; T03_descritivas_global.csv'))
    if num(falt.split('(')[1]) > 1:
        q.append(Q(f'Por que há {falt.split(" ")[0]} faltantes?',
                   'O setor fica sem valor quando o IBGE não divulgou o dado (sigilo, em setores com poucos '
                   'moradores) ou quando o denominador é zero. Não separei as duas causas por indicador. O setor '
                   'sai só deste indicador; na fatorial, sai inteiro (exclusão por lista).',
                   f'{sid}, linha Faltantes; T02_perdas_por_indicador.csv'))
    a = num(assim)
    if len(q) < 3:
        if abs(a) < 0.5:
            r = f'Quase simétrica (assimetria {assim}): média ({media}) e mediana ({med}) ficam perto.'
        elif a > 0:
            r = (f'Assimetria {assim}: cauda à direita. Poucos setores com valor bem alto puxam a média ({media}) '
                 f'para cima da mediana ({med}).')
        else:
            r = (f'Assimetria {assim}: cauda à esquerda. A maioria dos setores tem valor alto e poucos têm valor '
                 f'baixo; a média ({media}) fica abaixo da mediana ({med}).')
        q.append(Q('A distribuição é simétrica?', r, f'{sid}, linhas Assimetria, Média e Mediana'))
    if ivs7 == 'Sim' and ref.startswith('Sim'):
        r = ('Está no IVS-7 (os sete indicadores adaptados do IVS-BH 2012) e na referência proposta do NB05'
             + (', invertida: quanto menor a renda, maior a vulnerabilidade.' if 'invertida' in ref else '.'))
    elif ivs7 == 'Sim':
        r = ('Está no IVS-7, mas a referência proposta do NB05 é a versão sem lixo: com o lixo, ele forma um fator '
             'próprio (FP08).')
    elif ref.startswith('Sim'):
        r = ('Não estava no IVS-BH 2012. Entrou na referência proposta como a forma graduada do sem banheiro '
             '(demanda 6).')
    elif nome in POR_QUE_FORA:
        r = POR_QUE_FORA[nome]
    elif nome in MORADIA:
        r = 'Descreve o tipo de moradia (demanda 9). Na fatorial testei só a moradia não convencional.'
    else:
        r = ('Descreve a idade da população, contexto do ELSI (idosos). Não entra no índice de '
             'vulnerabilidade.')
    q.append(Q('Por que entra (ou não) no índice?', r,
               f'{sid}, linhas "Entra no IVS-7?" e "Entra na referência proposta?"; T20_variaveis_cargas.csv'))
    if num(cv) > 100 and len(q) < 4:
        q.append(Q(f'O que quer dizer o coeficiente de variação de {cv}?',
                   'O desvio-padrão é maior que a média: o valor muda muito de um setor para outro.',
                   f'{sid}, linha Coeficiente de variação'))
    k = [f'Setores com valor: {n}; faltantes: {falt}', f'Média {media}; mediana {med}',
         f'Setores com zero: {zeros}', f'IVS-7: {ivs7}; referência: {ref}']
    return dict(o=o, q=q[:4], k=k, ver=['DM01', 'VP03'])


def ds2(sid):
    base = sid.split('-')[0]
    nome = titulo(base)
    cab = next(ln for ln in V[sid]['linhas'] if ln[0] == 'Recorte')
    regs = [ln for ln in V[sid]['linhas'] if ln[0] in ('Norte', 'Nordeste', 'Sudeste', 'Sul', 'Centro-Oeste')]
    col = cab.index('% > 0') if '% > 0' in cab else cab.index('Mediana')
    med = cab[col] == 'Mediana'
    rot = 'mediana' if med else '% de setores com valor acima de zero'
    mx, mn = max(regs, key=lambda x: num(x[col])), min(regs, key=lambda x: num(x[col]))
    fcu, dem = cel(sid, 'Setores de FCU', col), cel(sid, 'Demais setores', col)
    o = (f'{nome} separado por região, por município e entre setores de favela (FCU) e os demais. ' +
         ('Compara a mediana de cada recorte (caixas sem os pontos extremos).' if med else
          'Como a maioria dos setores tem zero, compara o % de setores com valor acima de zero e a média.'))
    q = [Q('Onde é maior?', f'Na região {mx[0]} ({rot}: {mx[col]}); a menor é a {mn[0]} ({mn[col]}).',
           f'{sid}, tabela por recorte; T04_descritivas_regiao.csv'),
         Q('E nas favelas?', f'Nos setores de FCU, {fcu}; nos demais, {dem}. É descrição de grupos de setores, '
           'não de pessoas: num estudo ecológico não dá para passar para o indivíduo.',
           f'{sid}; ART008 p.194 (falácia ecológica); T15_fcu_por_indicador.csv')]
    m = re.search(r'Entre os municípios, (.+?)\.$', frase(sid))
    if m:
        q.append(Q('Qual município se destaca?', f'Entre os municípios, {m.group(1)}. A lista completa está '
                   'no anexo (aba Por_municipio).', f'{sid}; T05_descritivas_municipio.csv'))
    q.append(Q('Por que comparar com FCU?', 'FCU é a classificação do IBGE de Favelas e Comunidades Urbanas. '
               'Uso como critério externo: mais adiante, o índice é validado pela AUC contra FCU.',
               'ART005 p.46-47; FA14'))
    k = [f'{mx[0]}: {mx[col]}; {mn[0]}: {mn[col]}', f'FCU {fcu} × demais {dem}']
    return dict(o=o, q=q[:4], k=k)


# ---------------------------------------------------------------- DC: correlações
def dc(sid):
    f = frase(sid)
    o = ('Correlação de Spearman entre pares de indicadores, nos setores que têm os dois valores. Vai de -1 a 1: '
         'perto de 1, os setores ficam na mesma ordem nos dois; perto de -1, em ordem inversa; perto de 0, sem '
         'relação de ordem.')
    q = [Q('Por que Spearman e não Pearson?', 'Os indicadores são assimétricos e cheios de zeros. O Spearman usa '
           'a ordem (postos) dos setores e não depende de distribuição normal. O Pearson está no anexo.',
           'ART007 seção 11.2 (Guia FIOCRUZ); anexo, aba Correlacao_Pearson'),
         Q('Qual o par mais forte?', f, f'{sid}, matriz; T12_correlacao_spearman.csv')]
    extra = {
        'DC03': ('Por que Sem banheiro × Banheiro graduado dá 1,000?', 'O graduado é feito com as três partes do '
                 'sem banheiro, com pesos 1, 2 e 3. Os setores ficam praticamente na mesma ordem.'),
        'DC07': ('Por que Sem banheiro × Banheiro graduado dá 1,000?', 'O graduado é feito com as três partes do '
                 'sem banheiro, com pesos 1, 2 e 3. Os setores ficam praticamente na mesma ordem.'),
        'DC04': ('Por que as três rendas dão 1,00 entre si?', 'Elas só diferem no setor extremo de Belo '
                 'Horizonte; em todos os outros setores o valor é o mesmo.'),
        'DC06': ('Por que Pessoas 60+ e Dependência idosa andam juntas?', 'As duas têm os idosos no numerador; '
                 'muda só o denominador (toda a população ou só 15 a 59 anos).'),
    }
    if sid in extra:
        q.append(Q(extra[sid][0], extra[sid][1], f'{sid}; definições nos slides DS'))
    q.append(Q('O que fazer com dois indicadores muito correlacionados?', 'Na fatorial eles medem quase a mesma '
               'coisa. Por isso entra um ou outro (uma renda só, um tipo de banheiro, "não chega" ou "sem '
               'canalização"), nunca os dois.', 'VP03; FP04'))
    return dict(o=o, q=q[:4], k=[f], ver=['DS32'] if sid == 'DC06' else [])


# ---------------------------------------------------------------- FC: cenários
def fc(sid):
    tit = titulo(sid)
    sub = next(t for t in V[sid]['textos'] if 'variáveis ·' in t)
    n = re.search(r'([\d.]+) setores', sub).group(1)
    k_var = re.search(r'(\d+) variáveis', sub).group(1)
    saem = [ln for ln in V[sid]['linhas'] if len(ln) >= 3 and ln[0] not in ('Sai do S0 por causa de', 'Total que sai (alguma delas)', 'Variável', 'Região')
            and re.fullmatch(r'[\d.]+', ln[1] or '') and re.fullmatch(r'[\d.]*', ln[2] or '')]
    o = ('Um cenário é um conjunto de variáveis rodado numa fatorial. Entram os setores urbanos com dados '
         'divulgados e com todas as variáveis do cenário; por isso o número de setores muda de um cenário para '
         'outro. Tudo é comparado com o S0, que reproduz o IVS-7 do NB04.')
    try:
        total = cel(sid, 'Total que sai (alguma delas)')
    except KeyError:
        total = '0'
    if num(total) > 0:
        pior = max(saem, key=lambda x: num(x[1].replace('.', '')))
        r = (f'Entram os setores com as {k_var} variáveis: {n}. Saem {total} do S0; a variável que mais tira é '
             f'{pior[0]} ({pior[1]} setores sem o dado).')
    else:
        r = f'Entram os setores com as {k_var} variáveis: {n}. Nenhum setor do S0 sai.'
    q = [Q(f'Por que este cenário tem {n} setores?', r, f'{sid}, tabelas de faltantes e "Sai do S0"; cenarios.csv')]
    if sid == 'FC01':
        q.append(Q('Por que o S0 é a referência dos cenários?', 'Reproduz o IVS-7 do NB04, a fatorial que já '
                   'tínhamos. Os outros cenários mudam uma coisa por vez em relação a ele.', 'FC01; DM11'))
    if num(total) > 0:
        q.append(Q('Os setores que saem são diferentes dos que ficam?', 'Não comparei o perfil dos que saem neste '
                   'cenário; a tabela mostra só em que região ficam os que entram. Posso fazer essa comparação.',
                   f'{sid}, tabela por região [A CONFERIR: perfil dos que saem]'))
    q.append(Q('Por que não imputar os faltantes?', 'A fatorial precisa de todas as variáveis no setor. Imputar '
               'criaria valores que o IBGE não divulgou; preferi tirar o setor e mostrar quantos saem (exclusão '
               'por lista). É uma escolha que posso rever com você.', 'DS01; FA01; ART034 p.537 (limitações)'))
    return dict(o=o, q=q[:4], k=[f'{tit}', f'{n} setores; {k_var} variáveis; saem do S0: {total}'])


# ---------------------------------------------------------------- VP: um eixo por vez
def vp_eixo(sid):
    eixo = titulo(sid).split(': ')[-1]
    linhas = [ln for ln in V[sid]['linhas'] if '→' in ln[0]]
    o = ('Compara pares de versões iguais em tudo, menos neste eixo. Para cada troca: a mediana da diferença '
         'entre os pares e, entre colchetes, a menor e a maior. Spearman perto de 1 quer dizer que as duas '
         'versões ordenam os setores quase igual.')
    resumo = '; '.join(f'{ln[0]}: Spearman {ln[5].split(" ")[0]}, {ln[6].split(" ")[0]} mudam de faixa'
                       for ln in linhas[:2])
    q = [Q(f'Mudar {eixo} muda o índice?', resumo + '.', f'{sid}; pares_por_eixo.csv'),
         Q('E quantas versões passam nos critérios?', '; '.join(f'{ln[0]}: {ln[7]}' for ln in linhas[:3]) + '.',
           f'{sid}, coluna "Elegíveis antes → depois"'),
         Q('O que é "mudar de faixa"?', 'É o setor passar de uma faixa (baixo, médio, elevado, muito elevado) '
           'para outra entre as duas versões, pela regra do IVS-BH. Mostra o efeito prático da troca.',
           'FA13; ART003 p.10 (regra das faixas)')]
    return dict(o=o, q=q, k=[f'{ln[0]}: {ln[6].split(" ")[0]} mudam de faixa' for ln in linhas[:3]])


# ---------------------------------------------------------------- slides escritos à mão
M = {}
M['A01'] = dict(
    o='Capa. O relatório tem cinco partes: as demandas de setembro e como foram cobertas, a descrição de cada '
      'indicador, as correlações, a análise fatorial (a referência proposta e as demandas sobre ela) e as 108 '
      'versões da fatorial.',
    q=[Q('Qual o objetivo deste relatório?', 'Mostrar o que fiz de cada demanda de setembro, para você conferir '
         'os números com os seus resultados e decidir o que fica em aberto.', 'A01; DM01'),
       Q('Qual a base de dados?', 'Os setores censitários dos 70 municípios do ELSI-Brasil, Censo 2022: 109.032 '
         'setores, dos quais 104.108 são urbanos com dados divulgados.', 'DS01; T01_funil.csv')],
    k=['70 municípios do ELSI-Brasil; Censo 2022', '109.032 setores; 104.108 na descritiva; 73.878 na fatorial'],
    ver=['DS01', 'VP03'])
M['DM00'] = dict(
    o='Abre a parte das demandas de setembro. Cada slide tem: o pedido, como eu li, o que foi feito, o que '
      'contornei e onde está o resultado.',
    q=[Q('Todas as demandas foram atendidas?', 'As 11 estão concluídas. Três têm um ponto aberto: o banheiro '
         'graduado (demanda 6), a leitura de "fatorial de tudo isso" (demanda 8) e manter ou não a moradia não '
         'convencional (demanda 9).', 'DM01'),
       Q('O que é "o que contornamos"?', 'É onde o dado não deixou fazer exatamente o pedido, e o que fiz no '
         'lugar. Por exemplo, a mediana de renda do setor não existe no arquivo do Censo.', 'DM14; DM02')],
    k=['11 demandas, todas concluídas', 'Pontos abertos: demandas 6, 8 e 9'], ver=['DM01'])
M['DM01'] = dict(
    o='A lista das 11 demandas de setembro, o estado de cada uma, o slide que a explica e o slide do resultado.',
    q=[Q('A demanda 4 aparece com o DM05, que não está no relatório. Por quê?', 'Tirei o DM05 na edição. O '
         'resultado da demanda 4 (separar útil × inútil) está no mapa, como diz a coluna "Resultado em".', 'DM01, linha 4 [A CONFERIR: o Pedro confirma]'),
       Q('O que quer dizer "graduada (V00237) em aberto" na demanda 6?', 'O sem banheiro foi juntado com a '
         'V00495. A versão graduada, que usa a V00237, foi feita depois, no NB05; fica em aberto se ela substitui '
         'a V00495.', 'DM01; DM07'),
       Q('Por que a demanda 8 diz "leitura padrão da pergunta 6"?', '"Fatorial de tudo isso" pode ser lida de '
         'mais de um jeito. Usei a leitura de juntar as quatro variáveis novas ao IVS-7 (cenário S6); falta '
         'confirmar com você.', 'DM09')],
    k=['Demanda 6: V00495 feita; graduada em aberto', 'Demanda 9: moradia como não convencional, sem excluir'])
M['DM02'] = dict(
    o='Demanda 1: uma coluna de renda com a mediana. A renda do Censo por setor é a média (V06004); a mediana do '
      'setor (V06006) não vem no arquivo. O setor extremo de Belo Horizonte tem renda média de R$ 170.418,06.',
    q=[Q('Por que não usou a mediana do setor?', 'Ela não existe no arquivo do Censo: vêm V06001 a V06005, não '
         'V06006. Por isso troquei só o setor extremo pela mediana de BH calculada sem ele.', 'DM02; ART004'),
       Q('Quanto isso muda?', 'Só 1 setor fica diferente. A média de BH vai de R$ 4.649,57 a R$ 4.649,47 '
         'conforme a alternativa, e a média nos 70 fica em R$ 4.185,80.', 'DM02, tabela; renda_imputacao_alternativas.csv'),
       Q('Por que esse setor é extremo?', 'É o maior valor da base, muito acima dos outros, e ele sozinho puxava a '
         'média de BH. Não investiguei o motivo no setor (por exemplo, poucos domicílios de renda muito alta).',
         'DM02 [A CONFERIR: características do setor 310620005650366]')],
    k=['Setor 310620005650366, R$ 170.418,06', 'Imputado: R$ 3.058,235 (mediana de BH sem o setor)',
       '1 setor difere; base passa de 104 para 105 colunas'])
M['DM03'] = dict(
    o='Demanda 2: o quadro de indicadores. São 28 indicadores em quatro dimensões; 7 deles formam o IVS-7, '
      'adaptado do IVS-BH 2012.',
    q=[Q('Onde está a definição de cada indicador?', 'Um slide por indicador (DS03 a DS33): numerador, '
         'denominador, arquivo do Censo e a distribuição. O quadro também está no Quadro_Indicadores.csv da '
         'entrega.', 'DM03; DS03 a DS33'),
       Q('Por que 28 aqui e 31 slides de indicador?', 'Os 28 do quadro mais as duas variantes de renda e o '
         'banheiro graduado, que foram criados depois, nas demandas 1 e 6.', 'DM03; DM14 linha 2; DS02'),
       Q('O que é o IVS-7?', 'Os sete componentes do IVS-BH 2012 adaptados ao Censo 2022: água, esgoto, lixo, '
         'razão de moradores, analfabetismo, renda e cor/raça.', 'ART003; ART004; DC01')],
    k=['28 indicadores: 6 demografia, 7 morfologia, 10 saneamento, 5 socioeconômica', '7 no IVS-7'],
    ver=['DS02', 'DC01', 'DC07'])
M['DM04'] = dict(
    o='Demanda 3: fatorial sem bootstrap, sem e com rotação. Rotação é girar os eixos dos fatores para cada '
      'variável carregar forte num fator só, o que facilita dar nome aos fatores.',
    q=[Q('Por que tirou o bootstrap?', 'Foi o pedido. Para ter uma ideia da incerteza, rodei a fatorial tirando um '
         'município de cada vez (S0, S1 e S6).', 'DM04; estabilidade.csv'),
       Q('Qual a diferença entre Varimax e promax?', 'A Varimax mantém os fatores sem correlação (ortogonal). A '
         'promax deixa os fatores se correlacionarem (oblíqua); por isso tem Φ, a correlação entre eles.',
         'ART002 p.169; ART001 p.38; FA10'),
       Q('Os pesos mudam muito com a rotação?', 'No S0, o peso de F1 vai de 73,40% (promax, estrutura) a 75,89% '
         '(sem rotação). Como diz Figueiredo, as rotações costumam dar resultados parecidos quando o padrão é '
         'claro.', 'DM04, tabela; ART002 p.169')],
    k=['S0: F1 75,89% sem rotação; 74,00% Varimax; 74,69% promax padrão', 'Φ no S0: 0,151'],
    p=['Promax: artigo original (ART025) em andamento'])
M['DM06'] = dict(
    o='Demanda 5: comparar "água não chega" (a água não chega ao domicílio por canalização). Fiz duas leituras: '
      '(a) contra as outras categorias de água e (b) contra os indicadores do IVS-7.',
    q=[Q('O que mostra a comparação?', '"Não chega" se correlaciona mais com as outras faltas de água: 0,514 com '
         'sem água canalizada e 0,506 com água só no terreno. Com o IVS-7, a maior é com água inadequada (0,459).',
         'DM06, tabela; FP03'),
       Q('Por que o % de zeros mudou de 85,7% para 94,5%?', 'O 85,7% usava os 104.108 setores como denominador e '
         'contava o faltante como não zero. Sobre os setores com valor, são 94,5%.', 'DM06'),
       Q('Os setores com "não chega" são mais vulneráveis?', 'Na mediana, sim: renda menor e analfabetismo, '
         'cor/raça PPI e esgoto inadequado maiores (FP03). É descrição, não causa.', 'FP03')],
    k=['Spearman com água inadequada: 0,459', 'Zeros: 94,5% dos setores com valor'], ver=['FP03'])
M['DM07'] = dict(
    o='Demanda 6: juntar o "sem banheiro". A V00495 (sem banheiro de uso exclusivo) é a soma de três partes: '
      'V00236 (só banheiro de uso comum), V00237 (só sanitário) e V00238 (sem banheiro nem sanitário).',
    q=[Q('Como sabe que a V00495 contém as outras?', 'Conferi no banco: a soma V00236 + V00237 + V00238 fecha '
         'com a V00495 em 91.199 de 91.199 setores.', 'DM07, tabela'),
       Q('O que é o banheiro graduado?', 'Uma nota de 0 a 1 que pesa a gravidade: 1 para uso comum, 2 para só '
         'sanitário, 3 para sem banheiro nem sanitário, dividido por 3 vezes o total de domicílios.', 'DS13'),
       Q('Por que a V00237 não estava na base?', 'Não tinha sido extraída na primeira montagem. Extraí de novo dos '
         'arquivos do Censo.', 'DM07; DM14')],
    k=['V00495 = V00236 + V00237 + V00238 em 91.199 de 91.199', 'Graduado: pesos 1, 2 e 3'], ver=['DS13'])
M['DM08'] = dict(
    o='Demanda 7: acrescentar canalização e sem banheiro. Uma variável nova por vez sobre o IVS-7: não chega '
      '(S2), sem canalização (S3) e sem banheiro (S4).',
    q=[Q('Por que uma por vez?', 'Assim dá para ver o efeito de cada uma. Juntas, estão no S6 (demanda 8).',
         'DM08; DM09'),
       Q('Por que o S4 perde tantos setores?', 'O sem banheiro tem faltantes em muitos setores: o S4 fica com '
         '75.955, 11.590 a menos que o S0.', 'DM08, tabela; FC05'),
       Q('O KMO piora com as variáveis novas?', 'Não: fica entre 0,780 (S2) e 0,805 (S3); o S0 tem 0,783.',
         'DM08, tabela')],
    k=['S2: 78.256 setores; S3: 87.512; S4: 75.955', 'KMO de 0,780 a 0,805'], ver=['DM01'])
M['DM09'] = dict(
    o='Demanda 8: fatorial de tudo isso. Li como o IVS-7 mais as quatro variáveis novas (não chega, sem '
      'canalização, sem banheiro e não convencional): cenário S6. O S7 é o S6 sem lixo.',
    q=[Q('Por que o S6 explica menos variância?', 'Com 11 variáveis, dois fatores explicam 48,94%, contra 62,12% '
         'no IVS-7 (S0). Kaiser e Horn pedem 4 fatores no S6; fiquei com 2 para comparar com o IVS-BH.',
         'DM09, tabela; FP06'),
       Q('Essa leitura de "tudo isso" está certa?', 'É a minha leitura; preciso que você confirme. Se for outra, '
         'rodo de novo.', 'DM09'),
       Q('As três rendas mudam o resultado?', 'Não: o KMO é 0,810 nas três e a variância, 48,94%.', 'DM09, tabela')],
    k=['S6: 69.458 setores, KMO 0,810, 48,94%', 'S7: KMO 0,813, 53,22%'], ver=['FP06', 'FC07'])
M['DM10'] = dict(
    o='Demanda 9: habitação convencional. Usei no sentido da vulnerabilidade: moradia não convencional. Mostrei '
      'o diagnóstico em vez de tirar a variável.',
    q=[Q('Por que a moradia não convencional é problemática na fatorial?', 'Quase todo setor tem zero (92,19% '
         'no S5) e a comunalidade é muito baixa (0,013): a fatorial quase não "vê" a variável.', 'DM10, tabela'),
       Q('O que é comunalidade?', 'A parte da variância da variável que os fatores explicam. Abaixo de 0,30 '
         'ela fica fora pelos critérios do NB05.', 'ART002 p.176; FA04'),
       Q('Então tira ou mantém?', 'Pelo critério, sairia. Deixei o diagnóstico para decidirmos juntos.', 'DM10')],
    k=['Comunalidade 0,013 (S5); 0,018 (S6)', 'Zeros: 92,19% (S5)'])
M['DM11'] = dict(
    o='Demanda 10: testar lixo. Rodei com e sem o lixo: S1 (IVS-7 sem lixo) e S7 (S6 sem lixo).',
    q=[Q('O que acontece com o lixo?', 'No S0 ele forma um fator próprio na Varimax: Kaiser e Horn dão 2 '
         'fatores. Sem o lixo (S1) dão 1.', 'DM11, tabela; FP08'),
       Q('Por que o lixo se separa?', 'Quase não se correlaciona com os outros indicadores (0,095 com água '
         'inadequada). Ele sozinho vira um fator.', 'DC01; FP02'),
       Q('Isso é estável?', 'Não muito: em duas rodadas tirando um município de cada vez, o lixo muda de fator. E com '
         'lixo quase nenhuma versão passa nos critérios: só 3 das 108.', 'DM11; VP05')],
    k=['S0: 2 fatores (Kaiser e Horn); S1: 1', 'Lixo × água inadequada: 0,095'], ver=['DC01', 'VP05'])
M['DM12'] = dict(
    o='Demanda 11: testar as duas rotações. Varimax e promax; na promax há duas formas de ler as cargas: a '
      'matriz padrão (efeito próprio do fator) e a matriz estrutura (correlação da variável com o fator).',
    q=[Q('Qual rotação usar?', 'Não escolhi. Mostro as duas; os pesos de F1 no S1 ficam entre 59,64% e 65,82% '
         'nas leituras oblíquas e 65,01% na Varimax.', 'DM12, tabela; FP09'),
       Q('O que é o Φ de 0,522 no S1?', 'A correlação entre F1 e F2 na promax. Se fosse perto de zero, a Varimax '
         'bastaria.', 'DM12; FA10'),
       Q('Padrão ou estrutura?', 'A padrão mostra quanto cada fator, sozinho, explica a variável; a estrutura '
         'soma também o que vem pela correlação entre os fatores. Reportei as duas.', 'FA09')],
    k=['S1: F1 65,01% Varimax; 65,82% promax padrão; 59,64% estrutura', 'Φ no S1: 0,522'],
    p=['Matriz padrão × estrutura: sem leitura no fichamento'])
M['DM14'] = dict(
    o='Resumo de tudo o que contornei, demanda por demanda: o obstáculo, o que fiz no lugar e o efeito.',
    q=[Q('Algum contorno muda o resultado?', 'O maior efeito é a perda de setores quando entra variável nova '
         '(demanda 7). A renda imputada muda 1 setor.', 'DM14; DM08'),
       Q('O que é o erro dos 85,7%?', 'Um percentual calculado antes com o denominador errado; o certo é 94,5% '
         '(demanda 5).', 'DM06')],
    k=['Renda: 1 setor muda', 'Zeros de "não chega": 94,5%'])
M['DS00'] = dict(
    o='Abre a descritiva: um indicador por slide, sem cruzar indicadores. Cada um tem a definição, o arquivo do '
      'Censo e a distribuição nos 104.108 setores.',
    q=[Q('Por que 104.108 setores?', 'São os urbanos com Dados_sig = OK: tirei os rurais e os sem dados '
         'divulgados.', 'DS01'),
       Q('Por que não cruza indicadores aqui?', 'Foi a forma combinada: primeiro cada indicador sozinho; as '
         'correlações vêm depois (DC).', 'DC00')],
    k=['104.108 setores urbanos com Dados_sig = OK'])
M['DS01'] = dict(
    o='O funil: de 109.032 setores para os que entram em cada parte. É a forma recomendada de abrir os '
      'resultados (como se chegou à amostra).',
    q=[Q('Por que saem 2.685 setores rurais?', 'O IVS é de vulnerabilidade intraurbana: compara setores dentro da '
         'cidade, e os rurais ficam fora.', 'DS01; ART010 p.8'),
       Q('O que é Dados_sig?', 'Uma marca por setor: OK se o setor tem dados divulgados. Saem 2.239 setores sem '
         'dado (sigilo ou sem domicílio ocupado). A regra vem do IVS-BH.', 'ART004 (Tratamento dos dados); DS01'),
       Q('Por que a fatorial usa só 73.878?', 'Exclusão por lista: sai o setor que não tem algum dos sete '
         'componentes. Saem 30.230, a maior parte pelo analfabetismo e pelo banheiro graduado.', 'DS01; FA01')],
    k=['109.032 → 106.347 urbanos → 104.108 com dados → 73.878 na fatorial', 'Saem 30.230 na fatorial'],
    ver=['FA01'])
M['DS02'] = dict(
    o='Quantos setores têm valor em cada indicador. Faltante = o setor não tem o dado (sigilo ou denominador '
      'zero).',
    q=[Q('Qual indicador perde mais setores?', 'Analfabetismo 15+: 16.552 faltantes (15,90%). Depois água só no '
         'terreno (15.897) e banheiro graduado (14.292).', 'DS02, tabela; T02_perdas_por_indicador.csv'),
       Q('Por que o analfabetismo perde tanto?', 'Os dados de alfabetização têm mais setores sem divulgação. Não '
         'verifiquei a causa setor a setor.', 'DS02 [A CONFERIR: causa dos faltantes do analfabetismo]'),
       Q('Isso enviesa a fatorial?', 'Pode: o sigilo não é aleatório (atinge setores pequenos). Está nas '
         'limitações.', 'ART034 p.537; notas do Pedro (limitações)')],
    k=['Analfabetismo: 16.552 faltantes (15,90%)', 'Banheiro graduado: 14.292 (13,73%)'])
M['DC00'] = dict(
    o='Abre as correlações: primeiro os componentes do IVS-7, depois cada tema (saneamento, socioeconômica, '
      'morfologia, demografia) e por fim os 31 juntos.',
    q=[Q('Para que servem as correlações antes da fatorial?', 'A fatorial parte da matriz de correlação. Ver os '
         'pares antes mostra o que deve ficar junto num fator e o que é quase repetido.', 'FA02; ART002'),
       Q('Por que separar por tema?', 'Com os 31 juntos a matriz fica grande demais para ler; por tema dá para '
         'ver os pares de cada dimensão, e a matriz completa vem no fim (DC07).', 'DC07')],
    k=['Spearman em todas as matrizes'])

for s_ in ('DC01', 'DC02', 'DC03', 'DC04', 'DC05', 'DC06', 'DC07'):
    M[s_] = None  # modelo dc()

M['FA00'] = dict(
    o='Abre a fatorial da referência proposta do NB05: renda original, sem lixo, banheiro graduado, rotação '
      'promax, pesos 60/40 e normalização municipal. É uma proposta: a escolha é sua.',
    q=[Q('O que é análise fatorial?', 'Um método que resume muitas variáveis correlacionadas em poucos fatores '
         '(dimensões). Aqui gera duas dimensões, socioeconômica e saneamento, e a partir delas os pesos do índice.',
         'ART002 p.161; ART001'),
       Q('Por que essa é a referência?', 'Foi a primeira na ordem dos critérios fixados no NB05 (estabilidade, AUC '
         'no município e parcimônia) entre as elegíveis. É uma proposta para você avaliar, não uma decisão.',
         'estado_nb05 (referência proposta)'),
       Q('A fatorial é o caminho certo para os pesos?', 'Depende de o índice ser reflexivo ou formativo. Essa '
         'questão está em aberto e ainda vou ler os textos (Bollen; Diamantopoulos; Edwards).',
         'ART020; minhas notas da planilha')],
    k=['Renda original · sem lixo · banheiro graduado · promax · 60/40 · municipal'],
    p=['Formativo × reflexivo: ART020, ART021, ART022 planejados'])
M['FA01'] = dict(
    o='A amostra da fatorial: entram os setores com os sete componentes. Exclusão por lista: falta um, sai o '
      'setor.',
    q=[Q('Quem tira mais setores?', 'Analfabetismo 15+ (16.552 faltantes) e banheiro graduado (14.292).',
         'FA01, tabela'),
       Q('Por que não usar todos os 104.108?', 'A fatorial precisa de todas as variáveis no setor. Imputar criaria '
         'valores não divulgados; preferi tirar e mostrar quantos saem.', 'FA01; DS01')],
    k=['73.878 entram; 30.230 saem'])
M['FA02'] = dict(
    o='A matriz de correlação que entra na fatorial: Spearman entre os sete componentes, nos 73.878 setores.',
    q=[Q('Qual par é mais forte?', 'Renda (invertida) e cor/raça PPI: 0,784. Depois renda e analfabetismo: 0,753.',
         'FA02, matriz'),
       Q('Por que a renda está invertida?', 'Para todas as variáveis apontarem para o mesmo lado: valor maior = '
         'mais vulnerável. Renda alta é menos vulnerável, então entra com sinal trocado.', 'FA02; FA11'),
       Q('Por que fatorial com Spearman?', 'Os indicadores são assimétricos e cheios de zeros; a fatorial parte '
         'da matriz de postos, que não exige normalidade.', 'ART007 seção 11.2; FA02')],
    k=['Maior correlação: 0,784 (renda × cor/raça)', 'Menor: 0,083 (razão × banheiro graduado)'])
M['FA03'] = dict(
    o='Se os dados servem para fatorial. KMO compara as correlações entre as variáveis com as correlações '
      'parciais: vai de 0 a 1, e quanto maior, mais as variáveis têm algo em comum. Bartlett testa se a matriz '
      'de correlação é diferente da identidade.',
    q=[Q('O KMO de 0,7807 é bom?', 'Na escala de Matos (Enap) é "bom" (entre 0,7 e 0,8); na de Figueiredo, '
         '"mediano" (0,70 a 0,79). Passa o critério do NB05 (KMO ≥ 0,70).', 'ART001 p.43; ART002 p.166; FA16'),
       Q('E o Bartlett?', 'Rejeita a hipótese de identidade (qui-quadrado 185.197,02, 21 graus de liberdade). Mas '
         'com amostra grande o Bartlett quase sempre rejeita, então sozinho não basta.', 'ART001 p.43'),
       Q('Quanto os dois fatores explicam?', '60,82% da variância.', 'FA03; FA05')],
    k=['KMO 0,7807', 'n = 73.878; 7 variáveis', '2 fatores, 60,82%'])
M['FA04'] = dict(
    o='MSA é o KMO de cada variável: mostra se ela combina com as outras. Comunalidade é a parte da variância da '
      'variável que os dois fatores explicam.',
    q=[Q('Qual variável é a mais fraca?', 'Na MSA, renda (invertida), com 0,709; na comunalidade, razão de '
         'moradores, com 0,375. As duas passam os critérios (MSA ≥ 0,50; comunalidade ≥ 0,30).', 'FA04; FA16'),
       Q('Figueiredo pede comunalidade de 0,50. Por que 0,30?', 'O NB05 usou 0,30. Com 0,50, razão de moradores '
         '(0,375) e banheiro graduado (0,446) ficariam abaixo. É uma questão para decidirmos.', 'ART002 p.176; FA04'),
       Q('Por que a comunalidade oblíqua é igual?', 'Com dois fatores, a parte explicada da variável não muda com '
         'a rotação; muda só como se reparte entre F1 e F2.', 'FA04')],
    k=['Menor MSA: 0,709 (renda)', 'Menor comunalidade: 0,375 (razão de moradores)'])
M['FA05'] = dict(
    o='Autovalor é quanto da variância total um componente explica. Kaiser fica com os autovalores maiores que '
      '1. Horn (análise paralela) compara cada autovalor com o que daria em dados aleatórios do mesmo tamanho e '
      'fica com os que passam.',
    q=[Q('Quantos fatores?', 'Dois: os dois primeiros autovalores (3,186 e 1,071) passam de 1 e da referência de '
         'Horn (1,026 e 1,015).', 'FA05, tabela'),
       Q('O segundo fator não está no limite?', 'Está perto: 1,071 contra 1,015 de Horn. Por isso olho Kaiser e '
         'Horn juntos.', 'FA05'),
       Q('Kaiser é confiável aqui?', 'Matos diz que Kaiser é preciso com n > 250 e comunalidade média ≥ 0,6. O n '
         'passa; a comunalidade média fica no limite.', 'ART001 p.29')],
    k=['Autovalores: 3,186 e 1,071', 'Horn: 1,026 e 1,015', '60,82% acumulado'],
    p=['Horn: artigo original (ART024) em andamento; a definição está em Figueiredo, ART002 p.168'])
M['FA06'] = dict(
    o='A decisão do número de fatores: Kaiser e Horn concordam em 2, e a solução fica com 2.',
    q=[Q('Por que reter 2 e não 1?', 'Os dois critérios dão 2 nesta base. Também é o desenho do IVS-BH: uma '
         'dimensão socioeconômica e uma de saneamento.', 'FA06; ART003'),
       Q('E se desse 3?', 'Nos cenários com mais variáveis Kaiser e Horn pedem 3 ou 4; mantive 2 para comparar com '
         'o IVS-BH. Está no FP06.', 'FP06')],
    k=['Kaiser 2; Horn 2; retidos 2'], ver=['FP06'])
M['FA07'] = dict(
    o='Cargas sem rotação: a correlação de cada variável com cada fator antes de girar os eixos. O primeiro '
      'fator pega um pouco de tudo, o que dificulta dar nome.',
    q=[Q('O que é carga?', 'A correlação entre a variável e o fator. Acima de 0,40 em módulo, a variável é '
         'considerada parte do fator (critério do NB05).', 'FA16; ART002'),
       Q('Por que mostrar sem rotação?', 'Para comparar: a rotação não muda o total explicado, só reparte entre '
         'os fatores.', 'FA07; FA08')],
    k=['Maior carga em F1: renda (0,868)', 'Maior em F2: água inadequada (0,625)'])
M['FA08'] = dict(
    o='Cargas Varimax: rotação ortogonal, que deixa cada variável carregar forte num fator só. Os fatores ficam '
      'sem correlação entre si.',
    q=[Q('Como ficam os dois fatores?', 'F1 socioeconômico (renda 0,906, analfabetismo 0,861, cor/raça 0,848, '
         'razão de moradores 0,592); F2 saneamento (água 0,752, banheiro graduado 0,668, esgoto 0,560).',
         'FA08, tabela'),
       Q('E o esgoto?', 'Carrega nos dois (0,447 e 0,560): fica entre as dimensões.', 'FA08')],
    k=['F1: renda, analfabetismo, cor/raça, razão', 'F2: água, banheiro, esgoto'])
M['FA09'] = dict(
    o='Cargas promax, a rotação oblíqua da referência. Padrão: efeito próprio de cada fator. Estrutura: '
      'correlação total da variável com o fator, incluindo o que vem pela correlação entre os fatores.',
    q=[Q('Por que promax e não Varimax?', 'Os fatores são correlacionados (Φ = 0,411); a promax permite isso. '
         'Matos descreve a promax como um método rápido para bases grandes.', 'FA10; ART001 p.38'),
       Q('Qual matriz vira peso?', 'A padrão: os pesos saem do quadrado das cargas padrão.', 'FA11'),
       Q('Por que há carga negativa (-0,207 no banheiro em F1)?', 'Na matriz padrão, descontado o que vem de F2, '
         'o banheiro quase não tem relação própria com F1. É pequena e não entra no peso de F1.', 'FA09; FA11')],
    k=['Padrão: renda 0,937 em F1; água 0,783 em F2', 'Estrutura: renda 0,908; água 0,772'],
    p=['Promax: ART025 em andamento; padrão × estrutura sem leitura no fichamento'], ver=['FA10'])
M['FA10'] = dict(
    o='Φ é a correlação entre os dois fatores na rotação oblíqua. Na Varimax ela é zero por construção.',
    q=[Q('O Φ de 0,411 é alto?', 'É moderado: setores com pior condição socioeconômica tendem a ter pior '
         'saneamento. Por isso uma rotação oblíqua faz sentido.', 'FA10'),
       Q('E se fosse perto de zero?', 'A Varimax daria quase o mesmo e seria mais simples.', 'ART002 p.169')],
    k=['Φ = 0,411'])
M['FA11'] = dict(
    o='Como as cargas viram pesos: em cada dimensão, o peso da variável é o quadrado da sua carga padrão '
      'dividido pela soma dos quadrados; depois, a dimensão socioeconômica vale 60% e a de saneamento 40%.',
    q=[Q('Por que 60/40?', 'É a repartição do IVS-BH 2012 (60,4% e 39,6%), arredondada.', 'ART009 p.7; ART003'),
       Q('Por que o quadrado das cargas?', 'Para medir a importância relativa de cada variável no fator, como '
         'no ISU de Passarelli-Araujo.', 'ART009 p.7'),
       Q('Quem pesa mais?', 'Renda (invertida), 19,28%; depois analfabetismo (17,36%) e água (17,32%).', 'FA11')],
    k=['F1 60% · F2 40%', 'Renda 19,28%; analfabetismo 17,36%; água 17,32%'])
M['FA12'] = dict(
    o='O índice: cada variável é normalizada de 0 a 1 dentro do município (min-max) e somada com os pesos. O '
      'resultado vai de 0 a 1; maior = mais vulnerável.',
    q=[Q('O que é a normalização municipal?', 'Cada variável vira (valor - mínimo) / (máximo - mínimo) entre os '
         'setores do próprio município. Assim o índice compara setores dentro da cidade.', 'ART004 (Conversão de '
         'escala); estado_nb05'),
       Q('Como se distribui?', 'Mediana 0,3306; metade dos setores entre 0,2832 e 0,3757 (P25 e P75); máximo '
         '0,9120.', 'FA12, tabela'),
       Q('Dá para comparar índice entre municípios?', 'Não diretamente: a escala é de cada município. Para '
         'comparar entre cidades, a normalização global é a alternativa testada no NB05.', 'estado_nb05')],
    k=['Mediana 0,3306; média 0,3338', 'P10 0,2313; P90 0,4311'])
M['FA13'] = dict(
    o='As faixas por duas regras. IVS-BH: médio = média ± 0,5 desvio-padrão; baixo, abaixo; elevado, até 1,5 DP '
      'acima; muito elevado, acima disso. Quartis: um quarto dos setores em cada faixa.',
    q=[Q('Quantos setores estão em vulnerabilidade muito elevada?', 'Pela regra do IVS-BH, 3.933 (5,32%).',
         'FA13, tabela'),
       Q('Por que 1 setor fica "indefinido"?', 'Um município tem um setor só nesta amostra: não tem desvio-padrão, '
         'então a regra não se aplica.', 'estado_nb05'),
       Q('Qual regra usar?', 'A do IVS-BH mantém a comparação com 2012; quartis forçam faixas do mesmo tamanho. '
         'Fica para decidirmos.', 'ART003 p.10')],
    k=['Muito elevado: 3.933 (5,32%)', 'Médio: 31.052 (42,03%)'])
M['FA14'] = dict(
    o='Validação contra FCU. A curva ROC mostra quão bem o índice separa setores de favela (FCU) dos demais; a '
      'AUC é a área sob a curva: 0,5 é acaso, 1 é separação perfeita. A AUC é a chance de um setor FCU '
      'sorteado ter índice maior que um não FCU sorteado.',
    q=[Q('A AUC de 0,8309 é boa?', 'Separa bem: sorteando um setor FCU e um não FCU, a chance de o FCU ter '
         'índice maior é 0,8309.', 'FA14'),
       Q('Por que validar contra FCU?', 'É um critério externo do IBGE, independente do índice. Mas a FCU já usa '
         'condições de moradia, então há alguma circularidade com o saneamento.', 'ART005 p.47; estado_nb05'),
       Q('Por que a AUC municipal?', 'O índice é normalizado por município; a AUC dentro de cada município é a '
         'comparação justa. Uso a mediana dos 31 municípios com ao menos 30 setores de FCU e 30 fora.', 'FA15')],
    k=['AUC global 0,8309; mediana municipal 0,8186', '31 municípios na mediana'],
    p=['ROC/AUC: sem leitura própria no fichamento'])
M['FA15'] = dict(
    o='A AUC em cada município, só onde há ao menos 30 setores de FCU e 30 fora (para a AUC não depender de '
      'poucos setores).',
    q=[Q('Por que só 31 dos 41?', 'Nos outros há menos de 30 setores de FCU ou menos de 30 fora.',
         'FA15'),
       Q('Por que Duque de Caxias é tão baixo?', 'Lá a AUC é 0,6037: o índice separa pouco FCU dos demais. Não '
         'investiguei o motivo.', 'FA15 [A CONFERIR: Duque de Caxias]')],
    k=['Mediana 0,8186', 'Menor 0,6037 (Duque de Caxias); maior 0,9544 (Campinas)'])
M['FA16'] = dict(
    o='Os critérios de elegibilidade do NB05, fixados antes de rodar: KMO ≥ 0,70; MSA ≥ 0,50 em todas; '
      'comunalidade ≥ 0,30 em todas; maior carga ≥ 0,40 em todas.',
    q=[Q('A referência passa em tudo?', 'Sim. O ponto mais perto do limite é a comunalidade da razão de '
         'moradores (0,3755) e a menor carga máxima é a do esgoto (0,5245).', 'FA16'),
       Q('Por que fixar os critérios antes?', 'Para a escolha não depender de qual resultado parece melhor.',
         'FA16')],
    k=['Comunalidade mínima 0,3755 (razão de moradores)', 'Menor carga máxima 0,5245 (esgoto)'])
M['FP00'] = dict(
    o='Abre as demandas sobre a fatorial: um slide por pedido, na ordem das demandas. Aqui a comparação entre '
      'cenários é o próprio pedido.',
    q=[Q('Por que os cenários antes das demandas?', 'Cada demanda compara cenários; os slides FC mostram antes '
         'quais setores entram em cada um.', 'FC01 a FC09'),
       Q('Os cenários usam o mesmo método da referência?', 'Não: os cenários vêm da fatorial ampliada (antes do '
         'NB05), com pesos Varimax; a referência é do NB05, com promax e 60/40.', 'FP11; FA00')],
    k=['Cenários S0 a S7 e variantes'])
M['FC09'] = dict(
    o='As variantes: mesmos setores do cenário de origem, mudando a renda (sem extremo ou com mediana) ou a '
      'escala (postos dentro do município).',
    q=[Q('O que é "postos dentro do município"?', 'Em vez do valor, cada setor recebe a sua posição (posto) '
         'entre os setores do mesmo município. Tira o efeito de escala entre cidades.', 'FC09; FP11'),
       Q('Por que a variante sem extremo tem um setor a menos?', 'O setor extremo de BH fica vazio nela.',
         'FC09; DM02')],
    k=['S0 sem extremo: 87.544 setores', 'S6 sem extremo: 69.457'])
M['FP01'] = dict(
    o='Demanda 1 na fatorial: o efeito das três rendas.',
    q=[Q('A renda escolhida muda a fatorial?', 'Não: KMO 0,783, 62,12% de variância e peso de F1 74,00% nas '
         'três. A troca afeta um setor só.', 'FP01, tabela'),
       Q('Então para que três rendas?', 'Para mostrar que o setor extremo não distorce o resultado.', 'DM02')],
    k=['Três rendas: peso F1 74,00% em todas'])
M['FP02'] = dict(
    o='Demanda 3: as cargas do S0 (IVS-7) sem rotação, Varimax e promax lado a lado.',
    q=[Q('O que muda entre as rotações?', 'Pouco: o lixo carrega no F2 nas três (0,880 a 0,919); renda, '
         'analfabetismo e cor/raça ficam no F1.', 'FP02, tabela'),
       Q('Por que o lixo sozinho no F2?', 'É o fator próprio do lixo: ele quase não se correlaciona com os outros.',
         'FP08; DC01')],
    k=['Lixo em F2: 0,880 / 0,919 / 0,915'])
M['FP03'] = dict(
    o='Demanda 5: o perfil (medianas) dos setores com água não chega = 0 e > 0.',
    q=[Q('Como são os setores com "não chega" > 0?', 'Têm renda mediana menor (R$ 1.526,40 contra R$ 2.830,52), '
         'mais analfabetismo e mais cor/raça PPI.', 'FP03, tabela'),
       Q('Quantos setores têm "não chega" > 0?', '5.190, contra 89.212 com zero.', 'FP03')],
    k=['> 0: 5.190 setores; = 0: 89.212', 'Renda mediana R$ 1.526,40 × R$ 2.830,52'])
M['FP04'] = dict(
    o='Demanda 6 na fatorial: o S4 usa a V00495; no NB05 as três versões do banheiro entram na grade.',
    q=[Q('Qual versão do banheiro passa mais?', 'Com a V00495, 96 de 108 especificações são elegíveis; fora e '
         'graduado, 48 de 108 cada.', 'FP04'),
       Q('Então por que a referência usa o graduado?', 'O graduado gradua a gravidade e ficou na frente na ordem '
         'dos critérios do NB05. Com a V00495 há empate no topo em outra leitura; está em aberto.',
         'estado_nb05 (sensibilidade)')],
    k=['V00495: 96 de 108 elegíveis', 'S4: 75.955 setores'])
M['FP05'] = dict(
    o='Demanda 7: cada cenário acrescenta uma variável ao IVS-7 e perde os setores sem ela.',
    q=[Q('Alguma variável nova melhora a fatorial?', 'O sem canalização (S3) tem o maior KMO (0,805) e quase não '
         'perde setores. Mas com as novas, Kaiser pede 3 fatores.', 'FP05, tabela'),
       Q('Por que a variância de 2 fatores cai?', 'Com mais variáveis, dois fatores explicam uma parte menor do '
         'total: de 62,12% (S0) para 56,37% (S4).', 'FP05')],
    k=['S3: KMO 0,805; S4: 56,37%'])
M['FP06'] = dict(
    o='Demanda 8: tudo junto (S6), sem lixo (S7) e as variantes.',
    q=[Q('Por que ficar com 2 fatores se Kaiser e Horn pedem 4?', 'Para manter o desenho do IVS-BH (duas '
         'dimensões) e comparar os cenários. Com 4, o índice teria outra estrutura; posso rodar se você quiser.',
         'FP06; FA06'),
       Q('E com postos dentro do município?', 'O KMO cai para 0,769 e Kaiser e Horn pedem 3.', 'FP06')],
    k=['S6: Kaiser 4, Horn 4, 48,94%', 'S7: Kaiser 3, Horn 2'])
M['FP07'] = dict(
    o='Demanda 9 na fatorial: a moradia não convencional no S5 e no S6.',
    q=[Q('A variável muda a fatorial?', 'O KMO do S5 é igual ao do S0 (0,783), mas a variância de 2 fatores cai '
         'de 62,12% para 54,45%: a variável entra sem pertencer a nenhum fator.', 'FP07; DM10'),
       Q('Por que o S5 não perde setores?', 'A moradia não convencional não tem faltantes: o S5 fica com os mesmos '
         '87.545 setores do S0.', 'FP07; DS02')],
    k=['S5: KMO 0,783; 54,45%'])
M['FP08'] = dict(
    o='Demanda 10: o lixo. Com ele, Kaiser e Horn dão 2 fatores e o lixo forma um fator próprio; sem ele, 1.',
    q=[Q('Então o IVS-7 tem um fator só de lixo?', 'Na Varimax do S0, sim: o segundo fator é praticamente o lixo. '
         'Sem lixo (S1) sobra um fator.', 'FP08; FP02'),
       Q('No S6 o lixo também se separa?', 'Não: com as quatro variáveis novas, ele não forma fator próprio.',
         'FP08')],
    k=['S0: 2 fatores; S1: 1', 'S6: lixo não forma fator próprio'])
M['FP09'] = dict(
    o='Demanda 11: os pesos nas duas rotações, em S0, S1 e S6; na promax pelas duas matrizes.',
    q=[Q('Onde a rotação muda mais o peso?', 'No S1: o peso de F1 vai de 77,17% (sem rotação) a 59,64% (promax, '
         'estrutura).', 'FP09'),
       Q('Por que o Φ muda tanto entre cenários?', 'Depende de quais variáveis entram: 0,151 no S0, 0,522 no S1 e '
         '0,435 no S6.', 'FP09')],
    k=['Φ: 0,151 (S0), 0,522 (S1), 0,435 (S6)'])
M['FP11'] = dict(
    o='A AUC contra FCU em cada cenário: o índice com pesos Varimax, a renda invertida sozinha e a média simples '
      'dos postos das variáveis.',
    q=[Q('O índice separa FCU melhor que a renda?', 'Não: a renda invertida sozinha tem AUC maior em quase todos '
         'os cenários (0,8819 contra 0,8377 no S0). O índice só fica à frente em duas variantes com postos dentro '
         'do município (S0_postos_mun e S6_postos_mun).',
         'FP11, tabela'),
       Q('O que é a média simples de postos?', 'Um índice sem fatorial: cada variável vira posto e tira-se a média. '
         'Serve de comparação.', 'FP11')],
    k=['S0: índice 0,8377; renda 0,8819; postos 0,8684'])
M['FP13'] = dict(
    o='O que o índice acrescenta à renda. A renda sozinha separa FCU melhor; dentro de cada faixa de renda do '
      'município, o índice separa melhor.',
    q=[Q('Se a renda sozinha é melhor, para que o índice?', 'No geral a renda vence (0,8900 × 0,8309). Mas entre '
         'setores com renda parecida, o índice separa FCU melhor (0,7062 × 0,5900): ele acrescenta o que a renda '
         'não vê.', 'FP13, tabela; acrescimo_renda.csv'),
       Q('Somar o índice à renda ajuda?', 'Melhora o modelo (queda da deviance de 135,45; z de 17,48), mas a AUC '
         'quase não muda.', 'FP13; estado_nb05'),
       Q('Em quantas especificações a renda vence?', 'Em todas: em nenhuma das 192 o índice supera a renda na AUC '
         'global.', 'FP13')],
    k=['AUC global: índice 0,8309; renda 0,8900', 'Nos estratos de renda: índice 0,7062; renda 0,5900'])
M['VP00'] = dict(
    o='Abre as 108 versões: cinco escolhas (renda, lixo, água, banheiro, moradia) combinadas, todas com o mesmo '
      'método (promax, 60/40, normalização municipal).',
    q=[Q('Por que 108?', '3 rendas × 2 lixo × 3 água × 3 banheiro × 2 moradia = 108.', 'VP03'),
       Q('Para que tantas versões?', 'Para ver o efeito de cada escolha, uma por vez, mantendo o resto igual. É '
         'uma análise de sensibilidade.', 'ART039 p.775')],
    k=['108 versões; 24 passam'], ver=['VP03', 'VP05'])
M['VP02'] = dict(
    o='Como ler as versões: método fixo, cada versão na sua amostra própria, comparações na amostra comum.',
    q=[Q('O que é a amostra comum?', 'Os 67.981 setores que têm as variáveis de todas as 108. Comparar na mesma '
         'amostra evita que a diferença venha dos setores e não da versão.', 'VP02'),
       Q('Os critérios são os mesmos do NB05?', 'Sim: KMO ≥ 0,70; MSA ≥ 0,50; comunalidade ≥ 0,30; carga ≥ 0,40.',
         'VP02; FA16')],
    k=['Amostra comum: 67.981 setores'])
M['VP03'] = dict(
    o='Os cinco eixos e as opções de cada um. Seis variáveis ficam em todas as versões.',
    q=[Q('Por que "não chega" e "sem canalização" não entram juntas?', 'Uma contém a outra: entrariam duas vezes '
         'a mesma informação.', 'VP03'),
       Q('Por que essas variáveis fixas?', 'São os componentes do IVS-7 que não estão em discussão: água, esgoto, '
         'razão de moradores, analfabetismo, renda e cor/raça.', 'VP03')],
    k=['3 × 2 × 3 × 3 × 2 = 108', 'De 6 a 10 variáveis'])
M['VP04'] = dict(
    o='Quantos setores cada versão tem na amostra própria, por opção de cada eixo.',
    q=[Q('Qual opção tira mais setores?', 'Acrescentar "não chega": a mediana cai para 69.480 setores.', 'VP04'),
       Q('A renda muda a amostra?', 'Quase nada: só o setor extremo de BH (68.000 × 68.001).', 'VP04')],
    k=['Amostras de 68.000 a 87.545 setores'])
M['VP05'] = dict(
    o='As 108 numa grade: cada célula é uma versão, em verde as que passam, com o primeiro critério que falha nas '
      'outras.',
    q=[Q('Por que só 24 passam?', 'Quase todas as que falham, falham na comunalidade (alguma variável abaixo de '
         '0,30); algumas também na carga. Com lixo só 3 passam; sem lixo, 21.', 'VP05; versoes.csv'),
       Q('Que variável costuma ficar com comunalidade baixa?', 'Lixo e moradia não convencional quase não se '
         'correlacionam com o resto (DC01, DM10).', 'DC01; DM10 [A CONFERIR: variável que falha em cada versão]')],
    k=['24 de 108 passam; com lixo 3, sem lixo 21'])
M['VP06'] = dict(
    o='O KMO de cada versão, na mesma grade.',
    q=[Q('Alguma versão falha no KMO?', 'Não: o menor é 0,7653 (V048), acima de 0,70. O KMO não separa as '
         'versões; o que separa é a comunalidade.', 'VP06; VP05'),
       Q('Qual versão tem o maior KMO?', 'A V033 (0,8143).', 'VP06')],
    k=['KMO de 0,765 a 0,814'])
M['VP07'] = dict(
    o='A AUC no município de cada versão, na amostra comum.',
    q=[Q('A versão de maior AUC passa nos critérios?', 'Sim: a V024 (0,8642) está na lista das que passam.',
         'VP07; VP13'),
       Q('A AUC muda muito entre versões?', 'Pouco: de 0,812 a 0,864.', 'VP07')],
    k=['AUC no município de 0,812 a 0,864'], ver=['VP13'])
M['VP13'] = dict(
    o='As versões que passam nos critérios, em ordem de código, com setores, KMO, AUC e o Spearman com a V023 '
      '(a referência do NB05).',
    q=[Q('A referência do NB05 está entre elas?', 'Sim: é a V023 (renda original, sem lixo, banheiro graduado), '
         'com 73.878 setores e KMO 0,781.', 'VP13'),
       Q('As que passam dão índices parecidos?', 'O Spearman com a V023 mostra isso: perto de 1 quer dizer quase '
         'a mesma ordem de setores.', 'VP13')],
    k=['V023 = referência do NB05: 73.878 setores, KMO 0,781'])
M['VP13-2'] = dict(
    o='Continuação da lista das versões que passam.',
    q=[Q('Por que em ordem de código?', 'Para não sugerir uma escolha; a ordem pelos critérios está no VP14.',
         'VP14'),
       Q('O que quer dizer o Spearman com a V023?', 'Quanto a versão ordena os setores igual à referência do NB05: '
         'perto de 1, quase a mesma ordem.', 'VP13')],
    k=['24 versões no total'], ver=['VP13'])
M['VP14'] = dict(
    o='As elegíveis ordenadas pela posição média em três critérios: estabilidade (a maior variação do peso de '
      'uma variável quando se tira um município), AUC no município e parcimônia (menos variáveis). Empate: a '
      'mais próxima do IVS-BH 2012.',
    q=[Q('Então a melhor é a V025?', 'Não escolhi. A V025 tem a menor posição média (9,2) por esses critérios; a '
         'decisão é sua.', 'VP14'),
       Q('Por que a referência do NB05 (V023) não está no topo?', 'Ela tem a maior variação de peso da lista (6,94 '
         'p.p.); a posição média é 12,3. A ordem do NB05 era em outra grade (com rotações e pesos diferentes).',
         'VP14; estado_nb05'),
       Q('O que é parcimônia?', 'Preferir a versão com menos variáveis quando o resultado é parecido.', 'VP14')],
    k=['V025: posição média 9,2', 'V023: 12,3; variação de peso 6,94 p.p.'])

P00 = dict(
    titulo='Perguntas gerais',
    o='Perguntas que podem vir em qualquer momento da reunião: andamento, como conferi, ferramentas, limitações e '
      'próximos passos.',
    q=[Q('Em que pé está o trabalho?', 'As 11 demandas de setembro estão feitas. A descritiva cobre 31 indicadores '
         'e as correlações; a fatorial tem uma referência proposta e 108 versões testadas. Falta você escolher a '
         'versão e decidir os pontos em aberto.', 'DM01; DC07; VP00'),
       Q('Como você conferiu os números?', 'Recalculei os indicadores direto dos códigos V do Censo, por um '
         'caminho separado, e bateram com o banco e com as tabelas. Os números do relatório saem das tabelas de '
         'resultado.', 'auditoria geral de 08/10 (17 checagens)'),
       Q('Que ferramentas você usou?', 'Python para tudo (scripts que geram tabelas e slides). [Escreva aqui, com '
         'as suas palavras, como usou IA e como conferiu.]', '[o Pedro decide a resposta]'),
       Q('Quais as limitações?', 'Sigilo não aleatório, exclusão por lista, sem análise espacial e validação só '
         'contra FCU. E é estudo ecológico: vale para setores, não para pessoas.', 'notas do Pedro (ART034 p.537); '
         'ART008 p.194')],
    k=['11 demandas; 31 indicadores; 108 versões; 24 elegíveis', 'Em aberto: versão do índice, regra das faixas, '
       'formativo × reflexivo, moradia'],
    ver=['DM01', 'DC07', 'VP00', 'VP05'])

GLOSSARIO = [
    ('Setor censitário', 'A menor área do Censo, com algumas centenas de domicílios. É a unidade do índice, como '
     'no IVS-BH.', 'ART003 p.6'),
    ('Dados_sig', 'Marca do setor: OK se tem dados divulgados. A regra vem da nota do IVS-BH.', 'ART004'),
    ('Sigilo', 'O IBGE não divulga contagens de setores com poucos moradores, para não identificar ninguém. O '
     'valor fica vazio.', 'DS02'),
    ('FCU', 'Favelas e Comunidades Urbanas: classificação do IBGE de 2022 (antes "aglomerados subnormais"). Usada '
     'como critério externo de validação.', 'ART005 p.46'),
    ('Proporção × razão', 'Proporção vai de 0 a 1 (parte do total). Razão compara duas quantidades e pode passar '
     'de 1 (por exemplo, moradores por dormitório).', 'DS03; DS16'),
    ('Exclusão por lista', 'Na fatorial, o setor sem alguma das variáveis sai inteiro.', 'DS01; FA01'),
    ('Imputação', 'Pôr um valor no lugar de um dado. Só usei para o setor extremo de renda de BH (mediana do '
     'município).', 'DM02'),
    ('Assimetria', 'Mede se a distribuição tem cauda: positiva, cauda à direita (poucos valores altos); negativa, '
     'à esquerda.', 'slides DS'),
    ('Coeficiente de variação (CV)', 'Desvio-padrão dividido pela média. Acima de 100%, o valor varia muito '
     'entre setores.', 'slides DS'),
    ('P10, P25, P75, P90', 'Percentis: abaixo do P10 ficam 10% dos setores; abaixo do P90, 90%.', 'slides DS'),
    ('Spearman', 'Correlação pela ordem (postos) dos setores; não exige normalidade. Vai de -1 a 1.',
     'ART007 seção 11.2'),
    ('Pearson', 'Correlação pelos valores; sensível a extremos e a muitos zeros.', 'ART007'),
    ('Análise fatorial', 'Resume variáveis correlacionadas em poucos fatores (dimensões latentes).',
     'ART002 p.161'),
    ('Componente × fator', 'Componente vem da variância total (componentes principais); fator, só da parte '
     'comum. Aqui os autovalores são dos componentes.', 'ART001'),
    ('Autovalor', 'Quanto da variância total um componente explica. A soma dos autovalores é o número de '
     'variáveis.', 'FA05'),
    ('Critério de Kaiser', 'Fica com os componentes de autovalor maior que 1. Preciso com n > 250 e comunalidade '
     'média ≥ 0,6.', 'ART001 p.29'),
    ('Análise paralela de Horn', 'Compara cada autovalor com o de dados aleatórios do mesmo tamanho; fica com os '
     'que passam.', 'ART002 p.168'),
    ('KMO', 'Kaiser-Meyer-Olkin: de 0 a 1, se as variáveis têm correlação suficiente para a fatorial. Acima de '
     '0,7, bom (Matos).', 'ART001 p.43; ART002 p.166'),
    ('MSA', 'O KMO de cada variável. Abaixo de 0,5, a variável não combina com as outras.', 'FA04'),
    ('Bartlett', 'Testa se a matriz de correlação é a identidade. Com amostra grande quase sempre rejeita; não '
     'basta sozinho.', 'ART001 p.43'),
    ('Comunalidade', 'Parte da variância da variável explicada pelos fatores. O NB05 exige ≥ 0,30; Figueiredo '
     'sugere 0,50.', 'ART002 p.176'),
    ('Carga', 'Correlação da variável com o fator. O NB05 exige que a maior carga de cada variável seja ≥ 0,40.',
     'FA16'),
    ('Rotação', 'Girar os eixos para cada variável carregar forte num fator só. Não muda o total explicado.',
     'ART002 p.169'),
    ('Varimax', 'Rotação ortogonal: fatores sem correlação.', 'ART002 p.169'),
    ('Promax', 'Rotação oblíqua: permite fatores correlacionados. Rápida, boa para bases grandes.', 'ART001 p.38'),
    ('Φ (phi)', 'Correlação entre os fatores numa rotação oblíqua.', 'FA10'),
    ('Matriz padrão', 'Na oblíqua: o efeito próprio de cada fator sobre a variável. É a que vira peso.', 'FA09'),
    ('Matriz estrutura', 'Na oblíqua: a correlação total variável-fator, incluindo o que vem do outro fator.',
     'FA09'),
    ('Pesos do índice', 'Quadrado das cargas, normalizado em cada dimensão, vezes o peso da dimensão (60/40).',
     'FA11; ART009 p.7'),
    ('Normalização min-max', '(valor - mínimo) / (máximo - mínimo): põe cada variável entre 0 e 1.',
     'ART004 (Conversão de escala)'),
    ('Normalização municipal', 'O min-max feito dentro de cada município. A global usa todos os setores juntos.',
     'FA12'),
    ('Faixas do IVS-BH', 'Médio = média ± 0,5 DP; baixo abaixo; elevado até 1,5 DP acima; muito elevado acima.',
     'ART003 p.10'),
    ('Quartis', 'Cortes que deixam um quarto dos setores em cada faixa.', 'FA13'),
    ('Curva ROC', 'Para cada corte do índice, a proporção de FCU acima do corte contra a de não FCU acima.',
     'FA14'),
    ('AUC', 'Área sob a curva ROC: 0,5 é acaso; 1, separação perfeita. Chance de um FCU ter índice maior que um '
     'não FCU.', 'FA14'),
    ('AUC municipal', 'A AUC calculada dentro de cada município; uso a mediana dos que têm ao menos 30 setores de '
     'FCU e 30 fora.', 'FA15'),
    ('Média de postos', 'Índice sem fatorial: cada variável vira posto e tira-se a média. Serve de comparação.',
     'FP11'),
    ('Amostra comum', 'Os setores que têm as variáveis de todas as versões comparadas.', 'VP02'),
    ('Estabilidade', 'Quanto os pesos mudam quando se tira um município de cada vez.', 'VP14'),
    ('Parcimônia', 'Preferir menos variáveis quando o resultado é parecido.', 'VP14'),
    ('Formativo × reflexivo', 'Reflexivo: os indicadores refletem algo latente (fatorial serve). Formativo: os '
     'indicadores formam o índice (a fatorial pode não ser o caminho para os pesos). Em aberto.', 'ART020'),
    ('Falácia ecológica', 'Associação entre setores não vale automaticamente para pessoas.', 'ART008 p.194'),
]


def main() -> None:
    saida = {'P00': dict(P00, tipo='geral')}
    for sid in V:
        if re.fullmatch(r'DS(0[3-9]|[12]\d|3[0-3])', sid):
            c = ds(sid)
        elif re.fullmatch(r'DS\d{2}-2', sid):
            c = ds2(sid)
        elif re.fullmatch(r'FC0[1-8]', sid):
            c = fc(sid)
        elif sid in ('VP08', 'VP09', 'VP10', 'VP11', 'VP12'):
            c = vp_eixo(sid)
        elif sid in M and M[sid] is None:
            c = dc(sid)
        elif sid in M:
            c = M[sid]
        else:
            raise SystemExit(f'sem conteúdo de apoio para {sid}')
        saida[sid] = dict(titulo=titulo(sid), o=c['o'], q=c['q'], k=c['k'], ver=c.get('ver', []),
                          p=c.get('p', []), tipo='slide')
    por = 3
    for i in range(0, len(GLOSSARIO), por):
        gid = f'G{i // por + 1:02d}'
        saida[gid] = dict(titulo=f'Glossário ({i // por + 1}/{-(-len(GLOSSARIO) // por)})', tipo='glossario',
                          termos=[dict(t=t, d=d, de_onde=f) for t, d, f in GLOSSARIO[i:i + por]])
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(json.dumps(saida, ensure_ascii=False, indent=1), encoding='utf-8')
    n_q = sum(len(v.get('q', [])) for v in saida.values())
    print(f'{len(saida)} apoios ({len(V)} slides + P00 + {sum(k.startswith("G") for k in saida)} de glossário); '
          f'{n_q} perguntas; {len(GLOSSARIO)} termos → {SAIDA.relative_to(RAIZ)}')


if __name__ == '__main__':
    main()
