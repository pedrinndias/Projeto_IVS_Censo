"""Versão 3 do relatório descritivo (C-07): o arquivo que vai para a orientadora, e o PDF.

Parte da v2 (não regera nada) e grava ARQUIVOS NOVOS:
- as notas do apresentador (fontes, requisitos, origem de cada texto) vão para
  `relatorio_descritivo/dados/notas_fontes_v2.md`, só para o Pedro, e saem do .pptx;
- metadados sem o nome da biblioteca que montou o deck (assunto e empresa);
- texto alternativo das figuras: a legenda ou o título do slide, no lugar do caminho do arquivo;
- DS23: a frase "um em cada quatro tem mais de 100,00%" (impossível) é reescrita;
- DS16 (C-08): "Entra na referência proposta?" passa de "Não" a "Sim, como renda invertida";
- PDF exportado pelo LibreOffice (Cambria é trocada pela Caladea, de mesmas medidas).

    uv run --with python-pptx --with pypdf python scripts/preparar_entrega_relatorio.py [v2.pptx] [v3.pptx]
"""
from __future__ import annotations

import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

RAIZ = Path(__file__).resolve().parents[1]
PASTA = RAIZ / 'relatorio_descritivo'
ENTRADA = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v2.pptx'
SAIDA = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v3.pptx'
NOTAS = PASTA / 'dados' / 'notas_fontes_v2.md'
ID = re.compile(r'^(A|DM|DS|DC|FA|FC|FP|Z|VP)\d{2}(-\d+)?$')
TITULO = 'Relatório descritivo do andamento: IVS Censo 2022'
DS23 = (re.compile(r'Metade dos setores tem até 100,00%; um em cada quatro tem mais de 100,00%\. '
                   r'Entre o P10 e o P90, vai de ([\d,]+%) a 100,00%; ([\d,]+%) dos setores têm valor zero\.'),
        r'Pelo menos metade dos setores tem 100,00%; o P10 é \1, e \2 dos setores têm valor zero.')


def sid(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and ID.match(sh.text_frame.text.strip()):
            return sh.text_frame.text.strip()
    return '?'


def titulo(slide) -> str:
    """O título é o primeiro texto do alto do slide que não é o ID nem o rótulo em maiúsculas."""
    for sh in sorted((s for s in slide.shapes if s.has_text_frame), key=lambda s: s.top):
        t = sh.text_frame.text.strip()
        if t and not ID.match(t) and not t.isdigit() and t != t.upper() and Emu(sh.top).inches < 1.3:
            return t.split('\n')[0]
    return ''


def main() -> None:
    entrada = Path(sys.argv[1]) if len(sys.argv) > 1 else ENTRADA
    saida = Path(sys.argv[2]) if len(sys.argv) > 2 else SAIDA
    if saida.resolve() == entrada.resolve():
        sys.exit('A saída não pode ser o deck de entrada: o script grava sempre um arquivo novo.')
    prs = Presentation(entrada)
    guardadas, ds23, ds16, alt = [], 0, 0, 0
    for n, s in enumerate(prs.slides, 1):
        ident = sid(s)
        # 1. notas: guarda e esvazia
        if s.has_notes_slide:
            caixas = [sh for sh in s.notes_slide.shapes if sh.has_text_frame]
            texto = '\n'.join(sh.text_frame.text for sh in caixas).strip()
            if texto:
                guardadas.append(f'## {n} · {ident}\n\n{texto}\n')
            for sh in caixas:
                sh.text_frame.text = ''
        # 2. texto alternativo das figuras
        legenda = next((sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame
                        and sh.text_frame.text.strip().startswith('Figura ')), '')
        legenda = re.sub(r'^Figura \d+ — ', '', legenda).split('\n')[0]
        for sh in s.shapes:
            if sh.shape_type == 13:  # imagem
                sh._element.xpath('./p:nvPicPr/p:cNvPr')[0].set('descr', legenda or titulo(s))
                alt += 1
        # 3b. DS16 (C-08): a referência usa a renda média invertida (renda_inv na T20)
        if ident == 'DS16':
            for sh in s.shapes:
                if sh.has_table:
                    for row in sh.table.rows:
                        c = row.cells
                        if c[0].text.strip() == 'Entra na referência proposta?' and c[1].text.strip() == 'Não':
                            c[1].text_frame.paragraphs[0].runs[0].text = 'Sim, como renda invertida'
                            ds16 += 1
        # 3. DS23
        if ident == 'DS23':
            for sh in s.shapes:
                if sh.has_text_frame:
                    for par in sh.text_frame.paragraphs:
                        for r in par.runs:
                            novo, k = DS23[0].subn(DS23[1], r.text)
                            r.text, ds23 = novo, ds23 + k
    cp = prs.core_properties
    cp.title, cp.subject, cp.keywords, cp.comments, cp.category = TITULO, '', '', '', ''
    cp.author = cp.last_modified_by = 'Pedro Dias Soares'
    with tempfile.TemporaryDirectory() as tmp:
        bruto = Path(tmp) / 'bruto.pptx'
        prs.save(bruto)
        # 4. app.xml (empresa) não é exposto pelo python-pptx: regrava a parte no zip
        with zipfile.ZipFile(bruto) as zin, zipfile.ZipFile(saida, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                dados = zin.read(item.filename)
                if item.filename == 'docProps/app.xml':
                    dados = re.sub(rb'<Company>[^<]*</Company>', b'<Company></Company>', dados)
                zout.writestr(item, dados)
    NOTAS.write_text('# Notas do apresentador da v2 (fontes de cada slide), guardadas antes de sair do deck\n\n'
                     + '\n'.join(guardadas), encoding='utf-8')
    print(f'{len(prs.slides)} slides → {saida.name}; notas guardadas: {len(guardadas)} → {NOTAS.name}; '
          f'texto alternativo: {alt} figuras; DS23: {ds23} frase; DS16: {ds16} célula')
    if ds23 != 1 or ds16 != 1:
        sys.exit('DS23 ou DS16: o texto esperado não foi encontrado.')
    # 5. PDF
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(saida)],
                       check=True, capture_output=True)
        pdf = Path(tmp) / (saida.stem + '.pdf')
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(pdf))
        w.add_metadata({'/Title': TITULO, '/Author': 'Pedro Dias Soares', '/Creator': '', '/Producer': ''})
        destino = saida.with_suffix('.pdf')
        with open(destino, 'wb') as f:
            w.write(f)
    print(f'PDF: {len(PdfReader(destino).pages)} páginas → {destino.name}')


if __name__ == '__main__':
    main()
