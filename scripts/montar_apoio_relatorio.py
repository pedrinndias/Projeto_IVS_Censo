"""Monta a v4 do relatório descritivo (apoio do Pedro na reunião) a partir da v3 e do apoio.json.

Depois de cada slide da v3 entra um slide de apoio (faixa "APOIO · NÃO COMPARTILHAR"); o mesmo texto,
com a linha "De onde" de cada resposta e as leituras pendentes, vai para as notas do slide original e
do slide de apoio. P00 (perguntas gerais) abre o deck; o glossário (G01…) fecha. A v3 não muda: grava
um arquivo novo e o PDF.

Se um bloco não couber, sai a última pergunta (ou o último item de cor), e nunca diminui a fonte. O
que sai do slide continua nas notas; o script lista os cortes.

    uv run --with python-pptx --with pypdf python scripts/montar_apoio_relatorio.py
"""
from __future__ import annotations

import copy
import json
import math
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parents[1]
PASTA = RAIZ / 'relatorio_descritivo'
V3 = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v3.pptx'
APOIO = PASTA / 'apoio' / 'apoio.json'
SAIDA = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v4_APOIO.pptx'
ID = re.compile(r'^(A|DM|DS|DC|FA|FC|FP|Z|VP)\d{2}(-\d+)?$')
FONTE = 'Cambria'
TINTA, CINZA, PETROLEO = RGBColor(0x1A, 0x1A, 0x18), RGBColor(0x56, 0x53, 0x4C), RGBColor(0x1F, 0x4E, 0x4A)
FUNDO, BRANCO = RGBColor(0xEE, 0xF2, 0xF1), RGBColor(0xFF, 0xFF, 0xFF)
FAIXA = 'APOIO · NÃO COMPARTILHAR · '
TOPO, BASE = 1.25, 7.25          # área útil, em polegadas
FS = 16                          # corpo do apoio (o prompt pede ao menos 11)
cortes = []


def altura(texto: str, larg: float, fs: float) -> float:
    """Altura estimada (pol.): quebra de linha palavra a palavra, caractere médio de 0,47 em (medido na folha de contato)."""
    cpl = max(1, int(larg * 72 / (fs * 0.47)))
    linhas = 0
    for par in texto.split('\n'):
        n, atual = 1, 0
        for w in par.split():
            if atual and atual + 1 + len(w) > cpl:
                n, atual = n + 1, 0
            atual = (atual + 1 + len(w)) if atual else len(w)
            while atual > cpl:
                n, atual = n + 1, atual - cpl
        linhas += n
    return linhas * fs * 1.2 / 72


def caixa(slide, x, y, w, h, partes, fs, cor=TINTA):
    """partes: lista de parágrafos; cada um é lista de (texto, negrito) ou dict(runs, fs, cor, antes)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    for i, par in enumerate(partes):
        cfg = par if isinstance(par, dict) else dict(runs=par)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if cfg.get('antes'):
            p.space_before = Pt(cfg['antes'])
        for txt, neg in cfg['runs']:
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = FONTE, Pt(cfg.get('fs', fs)), neg
            r.font.color.rgb = cfg.get('cor', cor)
    return tb


def fundo_e_faixa(slide, ident, titulo):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = FUNDO
    faixa = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.42))
    faixa.fill.solid()
    faixa.fill.fore_color.rgb = PETROLEO
    faixa.line.fill.background()
    tf = faixa.text_frame
    tf.margin_left = Inches(0.6)
    r = tf.paragraphs[0].add_run()
    r.text = FAIXA + ident
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONTE, Pt(12), True, BRANCO
    caixa(slide, 0.6, 0.6, 12.1, 0.5, [[(titulo, True)]], 20)


def apoio_slide(slide, ident, a):
    """Cada coluna é uma caixa só, com parágrafos: o PowerPoint distribui as linhas e nada se sobrepõe."""
    fundo_e_faixa(slide, ident, f'{ident} · {a["titulo"]}' if ident != 'P00' else a['titulo'])
    lab = lambda t, antes=0: dict(runs=[(t, True)], fs=14, cor=PETROLEO, antes=antes)
    # esquerda: o que é + de cor
    xl, wl = 0.6, 4.55
    itens = list(a['k'])
    while True:
        h = (0.3 + altura(a['o'], wl, FS) + 0.25 + 0.3
             + sum(altura('· ' + k, wl, FS) for k in itens))
        if TOPO + h <= BASE or not itens:
            break
        cortes.append(f'{ident}: item de cor "{itens[-1][:40]}"')
        itens.pop()
    partes = [lab('O que é'), dict(runs=[(a['o'], False)], antes=3)]
    if itens:
        partes += [lab('Ter na ponta da língua', 16)] + [dict(runs=[('· ' + k, False)], antes=3) for k in itens]
    caixa(slide, xl, TOPO, wl, min(h, BASE - TOPO), partes, FS)
    # direita: perguntas
    xr, wr = 5.5, 7.23
    qs = list(a['q'])
    while True:
        h = 0.3 + sum(altura('P: ' + q['p'], wr, FS) + altura('R: ' + q['r'], wr, FS) + 0.16 for q in qs)
        if TOPO + h <= BASE or len(qs) <= 2:
            break
        cortes.append(f'{ident}: pergunta "{qs[-1]["p"][:50]}"')
        qs.pop()
    if TOPO + h > BASE:
        cortes.append(f'{ident}: DUAS PERGUNTAS NÃO CABEM')
    partes = [lab('Perguntas prováveis')]
    for q in qs:
        partes += [dict(runs=[('P: ', True), (q['p'], True)], antes=9),
                   dict(runs=[('R: ', True), (q['r'], False)], antes=2)]
    caixa(slide, xr, TOPO, wr, min(h, BASE - TOPO), partes, FS)
    return TOPO + h


def glossario_slide(slide, ident, a):
    fundo_e_faixa(slide, ident, a['titulo'])
    y = TOPO + 0.05
    for t in a['termos']:
        caixa(slide, 0.6, y, 12.1, 0.36, [[(t['t'], True)]], 17, PETROLEO)
        h = altura(t['d'], 12.1, 15)
        caixa(slide, 0.6, y + 0.4, 12.1, h, [[(t['d'], False)]], 15)
        y += 0.4 + h + 0.45
    return y


def texto_notas(ident, a):
    if a.get('tipo') == 'glossario':
        return '\n'.join(f'{t["t"]}: {t["d"]} (De onde: {t["de_onde"]})' for t in a['termos'])
    linhas = [f'APOIO {ident}', '', 'O QUE É', a['o'], '', 'PERGUNTAS']
    for q in a['q']:
        linhas += [f'P: {q["p"]}', f'R: {q["r"]}', f'De onde: {q["de_onde"]}', '']
    linhas += ['NA PONTA DA LÍNGUA'] + [f'· {k}' for k in a['k']]
    if a.get('p'):
        linhas += ['', 'LEITURA PENDENTE (estudar antes): ' + '; '.join(a['p'])]
    return '\n'.join(linhas)


def escrever_notas(slide, texto, molde):
    ns = slide.notes_slide
    caixas = [sh for sh in ns.shapes if sh.has_text_frame]
    if not caixas:  # o modelo de notas do pptxgenjs não traz caixa de texto: copia a de um slide da v3
        ns.shapes._spTree.append(copy.deepcopy(molde))
        caixas = [sh for sh in ns.shapes if sh.has_text_frame]
    caixas[0].text_frame.text = texto
    for sh in caixas[1:]:
        sh.text_frame.text = ''


def sid(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and ID.match(sh.text_frame.text.strip()):
            return sh.text_frame.text.strip()
    return '?'


def main() -> None:
    A = json.loads(APOIO.read_text(encoding='utf-8'))
    prs = Presentation(V3)
    originais = list(prs.slides)
    ids = [sid(s) for s in originais]
    falta = [i for i in ids if i not in A]
    if falta:
        sys.exit(f'sem apoio para: {falta}')
    molde = next(sh._element for sh in originais[0].notes_slide.shapes if sh.has_text_frame)
    lst = prs.slides._sldIdLst
    layout = originais[0].slide_layout

    def novo(depois_de=None, primeiro=False):
        s = prs.slides.add_slide(layout)
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        el = lst[-1]
        lst.remove(el)
        if primeiro:
            lst.insert(0, el)
        elif depois_de is not None:
            depois_de.addnext(el)
        else:
            lst.append(el)
        return s

    ids_xml = list(lst)
    for orig, ident, el in zip(originais, ids, ids_xml):
        s = novo(depois_de=el)
        apoio_slide(s, ident, A[ident])
        notas = texto_notas(ident, A[ident])
        escrever_notas(s, notas, molde)
        escrever_notas(orig, notas, molde)
    s = novo(primeiro=True)
    apoio_slide(s, 'P00', A['P00'])
    escrever_notas(s, texto_notas('P00', A['P00']), molde)
    for gid in sorted(k for k in A if re.fullmatch(r'G\d{2}', k)):
        s = novo()
        glossario_slide(s, gid, A[gid])
        escrever_notas(s, texto_notas(gid, A[gid]), molde)
    cp = prs.core_properties
    cp.title, cp.subject = 'Apoio da reunião: relatório descritivo IVS Censo 2022', 'apoio, não compartilhar'
    prs.save(SAIDA)
    print(f'{len(prs.slides)} slides → {SAIDA.name} ({len(originais)} originais + {len(originais)} de apoio + '
          f'P00 + {sum(1 for k in A if k.startswith("G"))} de glossário)')
    print(f'cortes por espaço ({len(cortes)}; o texto continua nas notas):')
    for c in cortes:
        print('  ', c)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(SAIDA)],
                       check=True, capture_output=True)
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(Path(tmp) / (SAIDA.stem + '.pdf')))
        w.add_metadata({'/Title': 'Apoio da reunião (não compartilhar)', '/Author': 'Pedro Dias Soares',
                        '/Creator': '', '/Producer': ''})
        with open(SAIDA.with_suffix('.pdf'), 'wb') as f:
            w.write(f)
    print(f'PDF: {len(PdfReader(SAIDA.with_suffix(".pdf")).pages)} páginas → {SAIDA.with_suffix(".pdf").name}')


if __name__ == '__main__':
    main()
