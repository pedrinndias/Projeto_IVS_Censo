"""Versão 2 do relatório descritivo: correções de texto e de voz (C-06), sem mudar nenhum número.

Parte do deck editado à mão (não regera nada) e grava um ARQUIVO NOVO. Cada regra vale só nos
slides indicados e tem o número de ocorrências esperado; se uma regra não acha o texto, o script
avisa. As trocas são feitas dentro do parágrafo, preservando a formatação do primeiro run.

O que muda:
- a orientadora deixa de aparecer em terceira pessoa (capa, DM09, DM10, FA09, FP11, selo FA07–FA16);
- caminhos de código, nomes de arquivo e de função saem do slide ("Calculado em" vai para as notas);
- frases quebradas: mediana zero nos DS, "supera ... em 0", "Menor" para correlação negativa,
  "saem 0", "Kaiser e Horn indicam 2 e 2";
- restos de edição: "da da", "excluído a (", linhas vazias do DM14, frase cortada;
- travessão nos títulos ("Demanda 1 — a renda" vira "Demanda 1: a renda").

    uv run --with python-pptx python scripts/revisar_texto_relatorio.py [entrada.pptx] [saida.pptx]
"""
from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

RAIZ = Path(__file__).resolve().parents[1]
PASTA = RAIZ / 'relatorio_descritivo'
ENTRADA = PASTA / 'Relatorio_Descritivo_IVS_2026-10_com_graficos_e_cenarios.pptx'
SAIDA = PASTA / 'Relatorio_Descritivo_IVS_2026-10_v2.pptx'
ID = re.compile(r'^(A|DM|DS|DC|FA|FC|FP|Z|VP)\d{2}(-\d+)?$')
SELO = 'Referência proposta no NB05'


def mediana_zero(m):
    p75, p90, zero = m.group(1), m.group(2), m.group(3)
    nulo = re.fullmatch(r'0,0+%?', p75.strip())
    if not nulo:
        return f'{zero} dos setores têm valor zero; um em cada quatro tem mais de {p75}, e um em cada dez, mais de {p90}.'
    if not re.fullmatch(r'0,0+%?', p90.strip()):
        return f'{zero} dos setores têm valor zero; um em cada dez tem mais de {p90}.'
    return f'{zero} dos setores têm valor zero.'


def indicam(m):
    n = m.group(1)
    return f'indicam {n} {"fator" if n == "1" else "fatores"} cada'


def autovalores(m):
    a, b = m.group(1), m.group(2)
    if a == b:
        return f'{a} autovalores passam de 1 (critério de Kaiser) e também ficam acima da referência de Horn'
    return m.group(0)


# (slides, padrão, troca, ocorrências esperadas); slides None = todos
REGRAS = [
    ({'A01'}, r'a análise fatorial — para conferência com os resultados da orientadora',
     'a análise fatorial, para conferência com os seus resultados', 1),
    ({'DM01'}, r'Redação do registro, estado e onde está cada uma', 'Estado de cada uma e onde está o resultado', 1),
    ({'DM01'}, r'^Demanda \(registro\)$', 'Demanda', 1),
    ({'DM02'}, r'Função renda_imputada_mediana_municipal em src/ivs_censo/renda\.py; coluna nova', 'Coluna nova', 1),
    ({'DM02'}, r'CSV de sensibilidade', 'tabela de sensibilidade', 1),
    ({'DM02'}, r'só no setor excluído a \(', 'só no setor excluído (', 1),
    ({'DM04'}, r' \(redação do registro\)', '', 1),
    ({'DM07', 'DM14'}, r'Saída: reextração no NB01, sem gravar o notebook\.',
     'Saída: nova extração da V00237 dos arquivos do Censo.', 2),
    ({'DM08', 'DM14'}, r', comparado com o S0 na coluna perdidos_vs_S0', ', comparado com o S0 pelos setores perdidos', 2),
    ({'DM09'}, r'A leitura de (.)tudo isso(.) é nossa e está registrada como pergunta à orientadora\.',
     r'A leitura de \1tudo isso\2 é nossa; falta confirmar se é essa.', 1),
    ({'DM10'}, r'MSA, comunalidade e % de zeros em cargas\.csv', 'MSA, comunalidade e % de zeros da variável', 1),
    ({'DM10'}, r'mostramos o diagnóstico; manter ou não é decisão da orientadora\.',
     'mostramos o diagnóstico para decidirmos se ela fica.', 1),
    ({'DM14'}, r'mostramos o diagnóstico;\s*$', 'mostramos o diagnóstico para decidirmos se ela fica.', 1),
    ({'DM11'}, r' \(lixo_muda\)', '', 1),
    ({'DM11', 'DM14'}, r'Nenhum contorno\. Registro: em S0', 'Nenhum contorno. Em S0', 2),
    ({'DM12'}, r', sem dizer que uma se afasta ou se aproxima da literatura', '', 1),
    ({'DM12'}, r'pesos\.csv com Varimax e promax \(padrão e estrutura\) e Φ, em todos os cenários; figura fa_plano_s6\.png\.',
     'Pesos da Varimax e da promax (padrão e estrutura) e Φ, em todos os cenários; figura do plano dos fatores no S6.', 1),
    ({'DM12', 'DM14'}, r'O posicionador de rótulos da figura veio do NB04; a figura foi conferida rótulo a rótulo antes de ir para slide\.',
     'Nenhum contorno de dado.', 2),
    ({'DM14'}, r'imputação pela mediana do município, Efeito', 'imputação pela mediana do município. Efeito', 1),
    ({'DM14'}, r'O quadro foi ampliado no NB05: TODOS_INDICADORES sobe de 26 para 28, com pct_so_sanitario e pct_banheiro_comum\.',
     'O quadro foi ampliado no NB05: de 26 para 28 indicadores, com só sanitário e banheiro de uso comum.', 1),
    (None, r'^numerador soma várias variáveis: sigilo parcial em uma delas entra como zero \(min_count=1\)',
     'o numerador soma várias variáveis: a que estiver em sigilo conta como zero', 12),
    (None, r'sem corte em (\[[^\]]+\]) — não é proporção', r'sem corte em \1 (não é proporção)', 5),
    (None, r'^(.+?) — (% de setores com valor acima de zero e média)$', r'\1: \2', 15),
    (None, r'^Demanda (\d+) — ', r'Demanda \1: ', 19),
    ({'FC01'}, r'· demanda —$', '· referência', 1),
    ({'FA09'}, r' \(pergunta 7 à orientadora\)', '', 1),
    ({'FP11'}, r'^Pedido derivado das perguntas à orientadora$', 'Pedido derivado das perguntas em aberto', 1),
    ({'FP13'}, r'Nas ([\d.]+) especificações avaliadas, o índice supera a renda sozinha na AUC global em 0\.',
     r'Em nenhuma das \1 especificações avaliadas o índice supera a renda sozinha na AUC global.', 1),
    ({'FP00'}, r'Demandas da da fatorial', 'Demandas da fatorial', 1),
    ({'FA05'}, r'(\d+) autovalores passam de 1 \(critério de Kaiser\); (\d+) ficam acima da referência de Horn', autovalores, 1),
    ({'FA06', 'FP06', 'FP08'}, r'indicam (\d+) e \1(?: fatores)?', indicam, 4),
    (None, r'Menor: ([^.]+?) \((-[\d,]+)\)', r'Correlação negativa mais forte: \1 (\2)', 6),
    (None, r'Em relação ao S0, saem 0: os que não têm alguma das variáveis novas\. Um setor pode faltar em mais de uma; '
           r'a coluna .só por ela. conta os que saem por uma única\.',
     'Nenhum setor do S0 sai: a variável nova não falta em nenhum deles.', 1),
    (None, r'Metade dos setores tem até 0,0+%?; um em cada quatro tem mais de ([^.;]+?)\. '
           r'Entre o P10 e o P90, vai de 0,0+%? a ([^;]+?); ([\d,]+%) dos setores têm valor zero\.', mediana_zero, 15),
    (None, r'renda sem_extremo', 'renda sem o setor extremo', 5),
    (None, r'renda mediana_mun', 'renda com a mediana do município', 5),
    ({'VP13', 'VP13-2'}, r'24 versões; lista em ordem de código, sem ordenar por resultado', '24 versões, em ordem de código', 2),
]


def sid(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and ID.match(sh.text_frame.text.strip()):
            return sh.text_frame.text.strip()
    return '?'


def paragrafos(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            yield from sh.text_frame.paragraphs
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield from c.text_frame.paragraphs


def trocar(par, padrao, troca) -> int:
    """Troca dentro do parágrafo; um trecho que cruza runs fica no primeiro run do trecho."""
    runs = par.runs
    texto = ''.join(r.text for r in runs)
    achados = list(re.finditer(padrao, texto))
    for m in reversed(achados):
        novo = m.expand(troca) if isinstance(troca, str) else troca(m)
        ini, fim, pos = m.start(), m.end(), 0
        limites = []
        for r in runs:
            limites.append((pos, pos + len(r.text)))
            pos += len(r.text)
        i = next(k for k, (a, b) in enumerate(limites) if a <= ini < b or (ini == b == fim and a == b))
        j = next(k for k, (a, b) in enumerate(limites) if a < fim <= b) if fim > ini else i
        a_i, a_j = limites[i][0], limites[j][0]
        resto = runs[j].text[fim - a_j:]
        runs[i].text = runs[i].text[:ini - a_i] + novo + (resto if j != i else runs[i].text[fim - a_i:])
        for k in range(i + 1, j + 1):
            runs[k].text = ''
    return len(achados)


def nota(slide, linha: str) -> None:
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        tf = next(sh.text_frame for sh in ns.shapes if sh.has_text_frame)
    tf.add_paragraph().text = linha


def main() -> None:
    entrada = Path(sys.argv[1]) if len(sys.argv) > 1 else ENTRADA
    saida = Path(sys.argv[2]) if len(sys.argv) > 2 else SAIDA
    if saida.resolve() == entrada.resolve():
        sys.exit('A saída não pode ser o deck de entrada: o script grava sempre um arquivo novo.')
    prs = Presentation(entrada)
    feitos = Counter()
    selos = calc = vazias = 0
    for s in prs.slides:
        ident = sid(s)
        for n, (ids, padrao, troca, _) in enumerate(REGRAS):
            if ids is None or ident in ids:
                for par in paragrafos(s):
                    feitos[n] += trocar(par, padrao, troca)
        for sh in list(s.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(SELO):
                sh._element.getparent().remove(sh._element)
                selos += 1
            if not sh.has_table:
                continue
            tbl = sh.table._tbl
            for tr, row in zip(list(tbl.tr_lst), list(sh.table.rows)):
                textos = [c.text.strip() for c in row.cells]
                apagar = False
                if textos[0] == 'Calculado em':
                    nota(s, 'Calculado em: ' + ' · '.join(textos[1:]))
                    apagar, calc = True, calc + 1
                elif ident == 'DM14' and not any(textos):
                    apagar, vazias = True, vazias + 1
                if apagar:
                    sh.height = Emu(sh.height - row.height)
                    tbl.remove(tr)
    prs.save(saida)
    print(f'{len(prs.slides)} slides → {saida.name}')
    print(f'selo FA07–FA16 removido: {selos}; "Calculado em" para as notas: {calc}; linhas vazias do DM14: {vazias}')
    falhas = [(n, feitos[n], r[3]) for n, r in enumerate(REGRAS) if feitos[n] != r[3]]
    for n, achou, esperado in falhas:
        print(f'  REGRA {n} achou {achou}, esperado {esperado}: {REGRAS[n][1][:70]}')
    print('regras: todas com a contagem esperada' if not falhas else f'regras fora da contagem: {len(falhas)}')


if __name__ == '__main__':
    main()
