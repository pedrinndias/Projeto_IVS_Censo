"""Insere slides novos do relatório descritivo no deck EDITADO pelo Pedro, sem regerar nada.

O deck do Pedro tem edições à mão que a origem não conhece. Os slides novos de uma rodada
(a lista `novos` de `relatorio_descritivo/dados/insercoes_<rodada>.json`) são gerados no deck
completo (`gerar_relatorio_descritivo.js --saida <tmp>`) e copiados daqui para dentro do deck
editado, cada um logo depois do slide que o antecede na ordem da origem; se esse slide foi
apagado pelo Pedro, vale o anterior que existir. O número impresso no pé é renumerado.
Grava um ARQUIVO NOVO; o deck do Pedro fica intocado.

    uv run --with python-pptx python scripts/inserir_no_relatorio.py \
        <deck_editado.pptx> <deck_completo_gerado.pptx> <insercoes.json> <saida.pptx>
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from anexar_versoes_ao_relatorio import copiar, sid  # noqa: E402


def main() -> None:
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    editado, completo, ins, saida = (Path(a) for a in sys.argv[1:])
    if saida.resolve() == editado.resolve():
        sys.exit('A saída não pode ser o deck editado.')
    regra = json.loads(ins.read_text(encoding='utf-8'))
    ordem, novos = regra['ordem'], regra['novos']
    prs, fonte = Presentation(editado), Presentation(completo)
    ja = {sid(s) for s in prs.slides}
    if ja & set(novos):
        sys.exit(f'O deck já tem slides desta rodada: {sorted(ja & set(novos))[:5]}')
    origem = {sid(s): s for s in fonte.slides}
    faltam = [n for n in novos if n not in origem]
    if faltam:
        sys.exit(f'O deck gerado não tem: {faltam[:5]}')
    lst = prs.slides._sldIdLst
    for n in novos:
        copiar(origem[n], prs)
        el = list(lst)[-1]
        presentes = {sid(s): i for i, s in enumerate(prs.slides)}
        k = ordem.index(n) - 1
        while k >= 0 and ordem[k] not in presentes:
            k -= 1
        lst.remove(el)
        if k < 0:
            lst.insert(0, el)
        else:
            list(lst)[presentes[ordem[k]]].addnext(el)
    renum = 0
    for pos, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and Emu(sh.top).inches > 6.8:
                run = sh.text_frame.paragraphs[0].runs[0]
                if run.text != str(pos):
                    run.text, renum = str(pos), renum + 1
    prs.save(saida)
    print(f'{len(ja)} slides + {len(novos)} novos = {len(prs.slides)} → {saida}; {renum} números de página corrigidos')


if __name__ == '__main__':
    main()
